"""Synthetic-only contract tests for bounded PMS export dialects, not live files."""
from __future__ import annotations

import builtins
import csv
import hashlib
import io
import json
import traceback
import warnings
import zipfile

import pytest

from plat_harness.ingest.pms_normalizer import (
    RentRollNormalizationError,
    normalize_rent_roll,
)


CANARIES = (
    "SYNTHETIC Alice Canary", "canary-person@example.invalid", "+1-212-555-0199",
    "SECRET-TENANT-ID", "PII-FREETEXT-CANARY", "PII-SHEET-CANARY",
)
HEADERS = {
    "yardi": ["Unit", "Status", "Unit Type", "Resident Name", "Phone", "Email", "Tenant ID", "Notes", "Rent", "Record Type"],
    "realpage": ["Unit #", "Occupancy Status", "Unit Category", "Resident", "Phone Number", "Email Address", "Tenant ID", "Notes", "Rent", "Row Type"],
    "entrata": ["Unit Number", "Occupancy", "Space Type", "Tenant Name", "Telephone", "E-mail", "Tenant ID", "Notes", "Rent", "Record Type"],
}
COUNTS = {"occupied": 2, "vacant": 1, "down": 1, "total": 4}
COMMERCIAL = {"occupied": 1, "vacant": 1, "down": 0, "total": 2}
NULL_COUNTS = dict.fromkeys(COUNTS)
ROOT_KEYS = {"version", "pms_type", "source_sha256", "residential_units", "commercial_units", "counts", "summary", "issues", "status"}


def unit(number="101", status="Current", kind="Residential", role="unit"):
    return [number, status, kind, *CANARIES[:3], CANARIES[3], CANARIES[4], "1234.5600", role]


def rows(pms="yardi", summary=True):
    result = [["Synthetic Rent Roll", CANARIES[4]], [], HEADERS[pms],
              unit(), unit("102", "Notice"), unit("103", "Vacant Rented"),
              HEADERS[pms], unit("104", "Down"), unit("C1", "Occupied", "Retail"),
              unit("C2", "Vacant", "Commercial")]
    if summary:
        result += [["Scope", "Occupied", "Vacant", "Down", "Total"],
                   ["Residential", 2, 1, 1, 4], ["Commercial", 1, 1, 0, 2]]
    return result


def csv_stream(values, binary=False):
    stream = io.StringIO(newline="")
    csv.writer(stream, lineterminator="\n").writerows(values)
    text = stream.getvalue()
    return io.BytesIO(text.encode("utf-8")) if binary else io.StringIO(text)


def xlsx_stream(values, extra_sheet=False):
    from openpyxl import Workbook
    book = Workbook()
    book.active.title = CANARIES[5]
    if extra_sheet:
        book.active.append([CANARIES[0], CANARIES[4]])
        sheet = book.create_sheet("Another PII title")
    else:
        sheet = book.active
    for row in values:
        sheet.append(row)
    stream = io.BytesIO()
    book.save(stream)
    book.close()
    stream.seek(0)
    return stream


def issue_codes(result):
    return {item["code"] for item in result["issues"]}


def assert_no_pii(result):
    encoded = json.dumps(result, allow_nan=False)
    for canary in CANARIES:
        assert canary not in encoded
    assert "Another PII title" not in encoded
    assert "1234.5600" not in encoded  # Money intentionally out of scope.


@pytest.mark.parametrize("pms", list(HEADERS))
@pytest.mark.parametrize("format_", ["csv_text", "csv_bytes", "xlsx"])
def test_vendor_dialects_exact_counts_reconciliation_and_redaction(pms, format_):
    values = rows(pms)
    stream = xlsx_stream(values) if format_ == "xlsx" else csv_stream(values, format_ == "csv_bytes")
    raw = stream.getvalue()
    raw = raw.encode("utf-8") if isinstance(raw, str) else raw
    result = normalize_rent_roll(stream, pms)
    assert set(result) == ROOT_KEYS
    assert result["version"] == "1.0"
    assert result["pms_type"] == pms
    assert result["source_sha256"] == hashlib.sha256(raw).hexdigest()
    assert result["status"] == "normalized_unvalidated"
    assert result["counts"] == {"residential": COUNTS, "commercial": COMMERCIAL}
    assert len(result["residential_units"]) == 4
    assert len(result["commercial_units"]) == 2
    for scope in ("residential", "commercial"):
        assert result["summary"][scope]["status"] == "reconciled"
        assert result["summary"][scope]["reported_counts"] == result["counts"][scope]
    first = result["residential_units"][0]
    assert set(first) == {"unit_id", "status", "unit_type", "tenant_name", "phone", "email", "evidence"}
    assert (first["tenant_name"], first["phone"], first["email"]) == ("[REDACTED]",) * 3
    cite = first["evidence"][0]["unit_id"]
    assert cite == {"source_sha256": result["source_sha256"], "sheet": 1, "row": 4, "row_end": 4, "column": 1}
    summary_cite = result["summary"]["residential"]["citations"][0]["total"]
    assert summary_cite["row"] == 12 and summary_cite["column"] == 5
    assert_no_pii(result)


@pytest.mark.parametrize("alias,canonical", [("YARDI", "yardi"), (" Yardi Voyager ", "yardi"), ("Voyager", "yardi"), ("RealPage OneSite", "realpage"), ("ONESITE", "realpage"), ("Entrata", "entrata")])
def test_pms_aliases(alias, canonical):
    assert normalize_rent_roll(csv_stream(rows(canonical)), alias)["pms_type"] == canonical


def test_summary_absent_is_not_reconciled_and_commercial_zero_has_known_basis():
    result = normalize_rent_roll(csv_stream([HEADERS["yardi"], unit()]), "yardi")
    assert result["summary"]["residential"]["status"] == "absent"
    assert result["summary"]["commercial"]["status"] == "absent"
    assert result["counts"]["residential"] == {"occupied": 1, "vacant": 0, "down": 0, "total": 1}
    assert result["counts"]["commercial"] == {"occupied": 0, "vacant": 0, "down": 0, "total": 0}
    assert result["status"] == "normalized_unvalidated"


@pytest.mark.parametrize("summary_row", [["Residential", 3, 0, 1, 4], ["Residential", 2, 1, 1, 999]])
def test_summary_mismatch_blocks_without_overwriting_observed_counts(summary_row):
    values = rows()
    values[11] = summary_row
    result = normalize_rent_roll(csv_stream(values), "yardi")
    assert result["summary"]["residential"]["status"] == "mismatch"
    assert result["status"] == "blocked"
    assert "SUMMARY_MISMATCH" in issue_codes(result)
    assert result["counts"]["residential"] == COUNTS


@pytest.mark.parametrize("bad", ["", "unknown", "Leased", "Occupied/Vacant", CANARIES[4]])
def test_unknown_or_missing_status_blocks_all_four_counts(bad):
    result = normalize_rent_roll(csv_stream([HEADERS["yardi"], unit(status=bad)]), "yardi")
    assert result["status"] == "blocked"
    assert result["counts"] == {"residential": NULL_COUNTS, "commercial": NULL_COUNTS}
    assert "UNSUPPORTED_STATUS" in issue_codes(result)
    assert_no_pii(result)


@pytest.mark.parametrize("bad", ["", "1BR", "mixed", CANARIES[4]])
def test_unknown_type_fails_closed(bad):
    result = normalize_rent_roll(csv_stream([HEADERS["yardi"], unit(kind=bad)]), "yardi")
    assert result["status"] == "blocked"
    assert result["counts"]["residential"] == NULL_COUNTS
    assert "UNSUPPORTED_UNIT_TYPE" in issue_codes(result)
    assert_no_pii(result)


@pytest.mark.parametrize("status,expected", [("Occupied", "occupied"), ("Current", "occupied"), ("Notice Rented", "occupied"), ("Notice Unrented", "occupied"), ("Vacant", "vacant"), ("Vacant Unrented", "vacant"), ("Vacant Rented", "vacant"), ("Down", "down"), ("Offline", "down"), ("Out of Service", "down")])
def test_supported_statuses(status, expected):
    result = normalize_rent_roll(csv_stream([HEADERS["yardi"], unit(status=status)]), "yardi")
    assert result["counts"]["residential"][expected] == 1
    assert result["counts"]["residential"]["total"] == 1


def test_duplicate_identical_unit_and_explicit_charge_cotenant_rows_count_once():
    values = [HEADERS["yardi"], unit(), unit(), unit(status="", kind="", role="charge"), unit(status="Current", role="co-tenant")]
    result = normalize_rent_roll(csv_stream(values), "yardi")
    assert result["status"] == "normalized_unvalidated"
    assert result["counts"]["residential"]["total"] == 1
    assert len(result["residential_units"][0]["evidence"]) == 4
    assert "DUPLICATE_UNIT_EVIDENCE" in issue_codes(result)
    assert "CONTINUATION_EVIDENCE" in issue_codes(result)
    assert_no_pii(result)


@pytest.mark.parametrize("change", [{"status": "Vacant"}, {"kind": "Retail"}, {"status": "Down", "role": "charge"}])
def test_conflicting_duplicate_blocks_and_excludes_conflicted_unit(change):
    result = normalize_rent_roll(csv_stream([HEADERS["yardi"], unit(), unit(**change)]), "yardi")
    assert "DUPLICATE_UNIT_CONFLICT" in issue_codes(result)
    assert result["status"] == "blocked"
    assert result["counts"]["residential"] == NULL_COUNTS
    assert result["residential_units"] == result["commercial_units"] == []


@pytest.mark.parametrize("values", [[unit(status="", kind="", role="charge")], [unit(), unit("", "", "", "charge")]])
def test_orphan_or_unidentified_continuation_is_not_assumed(values):
    result = normalize_rent_roll(csv_stream([HEADERS["yardi"], *values]), "yardi")
    assert result["status"] == "blocked"
    assert result["counts"]["residential"] == NULL_COUNTS


def test_future_applicant_sections_and_rows_not_counted():
    values = [HEADERS["yardi"], unit(), unit("102", "Applicant"), ["Future Residents"], HEADERS["yardi"], unit("103"), ["Applicants"], unit("104"), ["Current Residents"], HEADERS["yardi"], unit("105", "Vacant")]
    result = normalize_rent_roll(csv_stream(values), "yardi")
    assert result["counts"]["residential"] == {"occupied": 1, "vacant": 1, "down": 0, "total": 2}
    assert {u["unit_id"] for u in result["residential_units"]} == {"101", "105"}
    assert "FUTURE_OR_APPLICANT_EXCLUDED" in issue_codes(result)


def test_unscoped_report_total_blocks_instead_of_counting_or_silently_ignoring():
    result = normalize_rent_roll(csv_stream([HEADERS["yardi"], unit(), ["Grand Total", "99"]]), "yardi")
    assert result["status"] == "blocked"
    assert "UNSCOPED_REPORT_TOTAL" in issue_codes(result)
    assert len(result["residential_units"]) == 1


@pytest.mark.parametrize("bad", ["", "-1", "1.0", "=1+1", CANARIES[4]])
def test_invalid_summary_not_zero_filled(bad):
    values = rows()
    values[11][1] = bad
    result = normalize_rent_roll(csv_stream(values), "yardi")
    assert result["status"] == "blocked"
    assert result["summary"]["residential"]["status"] == "invalid"
    assert result["summary"]["residential"]["reported_counts"] is None
    assert_no_pii(result)


def test_conflicting_duplicate_summary_blocked():
    values = rows() + [["Residential", 1, 2, 1, 4]]
    result = normalize_rent_roll(csv_stream(values), "yardi")
    assert result["status"] == "blocked"
    assert "SUMMARY_CONFLICT" in issue_codes(result)
    assert result["summary"]["residential"]["status"] == "invalid"


@pytest.mark.parametrize("bad", [CANARIES[0], CANARIES[1], CANARIES[2], "2125550199", "Alice", "=1+1", "../101", "A 101", "A101\nPII", "12345"])
def test_unsafe_unit_identifiers_never_emitted(bad):
    result = normalize_rent_roll(csv_stream([HEADERS["yardi"], unit(number=bad)]), "yardi")
    assert result["status"] == "blocked"
    assert result["counts"]["residential"] == NULL_COUNTS
    assert result["residential_units"] == []
    assert "INVALID_UNIT_ID" in issue_codes(result)
    assert_no_pii(result)


def test_multiline_csv_citations_are_physical_lines_not_record_indices():
    values = [["Synthetic\nReport"], HEADERS["yardi"], unit(), unit("102", "Vacant")]
    values[2][3] = "Synthetic\nTenant"
    stream = csv_stream(values)
    result = normalize_rent_roll(stream, "yardi")
    assert result["residential_units"][0]["evidence"][0]["unit_id"]["row"] == 4
    assert result["residential_units"][0]["evidence"][0]["unit_id"]["row_end"] == 5
    assert result["residential_units"][1]["evidence"][0]["unit_id"]["row"] == 6


def test_xlsx_positional_sheet_ids_and_no_metadata_leak():
    stream = xlsx_stream(rows(), extra_sheet=True)
    stream.name = CANARIES[0] + ".xlsx"
    result = normalize_rent_roll(stream, "yardi")
    assert result["residential_units"][0]["evidence"][0]["unit_id"]["sheet"] == 2
    assert_no_pii(result)


def test_xlsx_formula_required_cell_is_not_evaluated():
    values = rows()
    values[3][1] = '=HYPERLINK("https://example.invalid/PII","Occupied")'
    result = normalize_rent_roll(xlsx_stream(values), "yardi")
    assert result["status"] == "blocked"
    assert "UNSUPPORTED_STATUS" in issue_codes(result)
    assert result["counts"]["residential"] == NULL_COUNTS
    assert "HYPERLINK" not in json.dumps(result)


def test_bytes_bom_hash_exact_and_citations_change_when_source_changes():
    original = csv_stream(rows(), binary=True).getvalue()
    with_bom = b"\xef\xbb\xbf" + original
    first = normalize_rent_roll(io.BytesIO(original), "yardi")
    second = normalize_rent_roll(io.BytesIO(with_bom), "yardi")
    assert first["source_sha256"] != second["source_sha256"]
    assert second["source_sha256"] == hashlib.sha256(with_bom).hexdigest()
    assert second["residential_units"][0]["evidence"][0]["status"]["source_sha256"] == second["source_sha256"]


@pytest.mark.parametrize("stream,pms,code", [
    (io.StringIO(""), "yardi", "EMPTY_INPUT"),
    (io.StringIO("  \n"), "yardi", "EMPTY_INPUT"),
    (io.StringIO(CANARIES[4]), "yardi", "HEADER_NOT_FOUND"),
    (io.StringIO("Unit,Status\n101,Current\n"), "yardi", "HEADER_NOT_FOUND"),
    (io.StringIO("Unit,Status,Unit Type,Status\n101,Current,Residential,Vacant\n"), "yardi", "AMBIGUOUS_HEADER"),
    (io.StringIO("anything"), CANARIES[4], "UNSUPPORTED_PMS"),
    (io.StringIO("anything"), None, "UNSUPPORTED_PMS"),
    (io.BytesIO(b"\xff\xff"), "yardi", "MALFORMED_INPUT"),
    (io.BytesIO(b"PK\x03\x04not-a-workbook"), "yardi", "MALFORMED_INPUT"),
    (io.StringIO('Unit,Status,Unit Type\n101,"unclosed'), "yardi", "MALFORMED_INPUT"),
])
def test_typed_safe_errors(stream, pms, code):
    with pytest.raises(RentRollNormalizationError) as error:
        normalize_rent_roll(stream, pms)
    assert error.value.code == code
    assert error.value.__context__ is None
    assert_no_pii({"repr": repr(error.value), "error": str(error.value), "trace": "".join(traceback.format_exception(error.value))})


def test_stream_read_exception_is_sanitized_without_exception_chaining():
    class Broken:
        def read(self, size):
            raise RuntimeError(CANARIES[0])
    with pytest.raises(RentRollNormalizationError) as error:
        normalize_rent_roll(Broken(), "yardi")
    assert error.value.code == "MALFORMED_INPUT"
    assert error.value.__context__ is None
    assert CANARIES[0] not in "".join(traceback.format_exception(error.value))


def test_empty_inventory_cannot_produce_zero_denominator():
    result = normalize_rent_roll(csv_stream([HEADERS["yardi"]]), "yardi")
    assert result["status"] == "blocked"
    assert result["counts"] == {"residential": NULL_COUNTS, "commercial": NULL_COUNTS}
    assert "NO_CURRENT_UNITS" in issue_codes(result)


def test_unknown_record_role_and_ragged_rows_block():
    for bad_row in [unit(role=CANARIES[4]), ["101"], unit() + [CANARIES[4]]]:
        result = normalize_rent_roll(csv_stream([HEADERS["yardi"], bad_row]), "yardi")
        assert result["status"] == "blocked"
        assert result["counts"]["residential"] == NULL_COUNTS
        assert_no_pii(result)


def test_openpyxl_is_optional_lazy_and_unavailable_error_is_typed(monkeypatch):
    stream = xlsx_stream(rows())
    original_import = builtins.__import__
    def no_openpyxl(name, *args, **kwargs):
        if name == "openpyxl" or name.startswith("openpyxl."):
            raise ImportError(CANARIES[4])
        return original_import(name, *args, **kwargs)
    monkeypatch.setattr(builtins, "__import__", no_openpyxl)
    assert normalize_rent_roll(csv_stream(rows()), "yardi")["status"] == "normalized_unvalidated"
    with pytest.raises(RentRollNormalizationError) as error:
        normalize_rent_roll(stream, "yardi")
    assert error.value.code == "XLSX_DEPENDENCY_MISSING"
    assert error.value.__context__ is None
    assert_no_pii({"error": str(error.value), "repr": repr(error.value)})


def test_library_warning_cannot_leak_workbook_pii(monkeypatch):
    import openpyxl
    stream = xlsx_stream(rows())
    original_load = openpyxl.load_workbook
    def noisy_load(*args, **kwargs):
        warnings.warn(CANARIES[0], UserWarning)
        return original_load(*args, **kwargs)
    monkeypatch.setattr(openpyxl, "load_workbook", noisy_load)
    with warnings.catch_warnings(record=True) as emitted:
        warnings.simplefilter("always")
        with pytest.raises(RentRollNormalizationError) as error:
            normalize_rent_roll(stream, "yardi")
    assert emitted == []
    assert error.value.code == "MALFORMED_INPUT"
    assert error.value.__context__ is None
    assert_no_pii({"trace": "".join(traceback.format_exception(error.value))})


def test_input_byte_limit_is_typed_and_safe():
    with pytest.raises(RentRollNormalizationError) as error:
        normalize_rent_roll(io.BytesIO(b"x" * (8 * 1024 * 1024 + 1)), "yardi")
    assert error.value.code == "INPUT_LIMIT_EXCEEDED"


def test_xlsx_declared_dimensions_cannot_hide_rows():
    stream = xlsx_stream([HEADERS["yardi"], unit(), unit("102", "Vacant")])
    import re
    modified = io.BytesIO()
    with zipfile.ZipFile(stream) as original, zipfile.ZipFile(modified, "w") as output:
        for item in original.infolist():
            value = original.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                value = re.sub(rb'<dimension ref="[^"]+"', b'<dimension ref="A1:A1"', value)
            output.writestr(item, value)
    modified.seek(0)
    result = normalize_rent_roll(modified, "yardi")
    assert result["counts"]["residential"]["total"] == 2


def test_xlsx_macro_archive_is_rejected_without_echo():
    stream = xlsx_stream(rows())
    modified = io.BytesIO()
    with zipfile.ZipFile(stream) as original, zipfile.ZipFile(modified, "w") as output:
        for item in original.infolist():
            output.writestr(item, original.read(item.filename))
        output.writestr("xl/vbaProject.bin", CANARIES[0])
    modified.seek(0)
    with pytest.raises(RentRollNormalizationError) as error:
        normalize_rent_roll(modified, "yardi")
    assert error.value.code == "MALFORMED_INPUT"
    assert_no_pii({"error": str(error.value)})


def test_inventory_error_keeps_summary_unresolved_not_reconciled():
    values = rows()
    values[3][1] = ""
    result = normalize_rent_roll(csv_stream(values), "yardi")
    assert result["summary"]["residential"]["status"] == "unresolved"
    assert result["summary"]["residential"]["reported_counts"] == COUNTS
    assert result["counts"]["residential"] == NULL_COUNTS


def test_invalid_summary_cannot_be_healed_by_later_good_row():
    values = rows()
    values[11][1] = ""
    values += [["Residential", 2, 1, 1, 4]]
    result = normalize_rent_roll(csv_stream(values), "yardi")
    assert result["summary"]["residential"]["status"] == "invalid"
    assert result["summary"]["residential"]["reported_counts"] is None


def test_identical_summary_retains_all_citations_without_double_counting():
    result = normalize_rent_roll(csv_stream(rows() + [["Residential", 2, 1, 1, 4]]), "yardi")
    assert result["summary"]["residential"]["status"] == "reconciled"
    assert len(result["summary"]["residential"]["citations"]) == 2
    assert result["counts"]["residential"] == COUNTS


def test_commercial_only_is_not_fabricated_residential_occupancy():
    result = normalize_rent_roll(csv_stream([HEADERS["yardi"], unit("R1", "Down", "Retail")]), "yardi")
    assert result["commercial_units"][0]["unit_type"] == "commercial"
    assert result["residential_units"] == []
    assert result["counts"]["commercial"] == {"occupied": 0, "vacant": 0, "down": 1, "total": 1}
    assert result["counts"]["residential"] == {"occupied": 0, "vacant": 0, "down": 0, "total": 0}


def test_issue_objects_allowlist_static_codes_and_citations_only():
    result = normalize_rent_roll(csv_stream([HEADERS["yardi"], unit(status=CANARIES[4])]), "yardi")
    for issue in result["issues"]:
        assert set(issue) == {"code", "severity", "citation"}
        assert issue["severity"] in {"warning", "blocker"}
    assert_no_pii(result)
