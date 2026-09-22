"""Synthetic-only contract tests for Abilene split-header/lease-charge XLSX adapter.

All test workbooks are generated in-memory via openpyxl. No live or private deal files.
"""
from __future__ import annotations

import hashlib
import io
import json
import traceback
import zipfile

import pytest
from openpyxl import Workbook

from plat_harness.ingest.lease_charge_xlsx import (
    BASE_RENT_CODES,
    ANCILLARY_CATEGORIES,
    ANCILLARY_CODES,
    classify_charge,
    is_base_rent,
    is_ancillary_charge,
    normalize_lease_charge_xlsx,
)
from plat_harness.ingest.pms_normalizer import RentRollNormalizationError

CANARIES = (
    "SYNTHETIC CANARY TENANT",
    "John Doe Canary",
    "Jane Smith Canary",
    "canary_user@example.invalid",
    "+1-555-019-9999",
    "SECRET_PII_NOTE",
)

ROOT_KEYS = {
    "version",
    "pms_type",
    "source_sha256",
    "residential_units",
    "commercial_units",
    "counts",
    "summary",
    "issues",
    "status",
}


def build_synthetic_workbook(
    *,
    title: str = "Rent Roll with Lease Charges",
    row5_headers: list[str] | None = None,
    row6_headers: list[str] | None = None,
    section_header: str = "Current/Notice/Vacant Residents",
    unit_rows: list[list] | None = None,
    extra_sections: list[tuple[str, list[list]]] | None = None,
    totals_row: list | None = None,
) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Rent Roll"

    # Row 1: Title
    ws.cell(row=1, column=1, value=title)
    # Rows 2-4: Subheaders / metadata
    ws.cell(row=2, column=1, value="Property: Synthetic Asset")
    ws.cell(row=3, column=1, value="As of Date: 01/31/2026")
    ws.cell(row=4, column=1, value="")

    # Row 5 & 6: Split headers
    if row5_headers is None:
        row5_headers = [
            "Unit", "Unit Type", "Unit", "Resident Name", "Market",
            "Charge", "Amount", "Resident", "Move In", "Lease",
            "Move Out", "Balance",
        ]
    if row6_headers is None:
        row6_headers = [
            "", "", "Sq Ft", "", "Rent",
            "Code", "", "Deposit", "", "Expiration",
            "", "",
        ]

    for col_idx, val in enumerate(row5_headers, start=1):
        ws.cell(row=5, column=col_idx, value=val)
    for col_idx, val in enumerate(row6_headers, start=1):
        ws.cell(row=6, column=col_idx, value=val)

    # Row 7: Section header
    ws.cell(row=7, column=1, value=section_header)

    # Row 8+: Unit rows
    current_row = 8
    if unit_rows is not None:
        for r_vals in unit_rows:
            for col_idx, val in enumerate(r_vals, start=1):
                ws.cell(row=current_row, column=col_idx, value=val)
            current_row += 1

    # Extra sections (e.g. Future Residents, Applicants)
    if extra_sections:
        for sec_name, sec_rows in extra_sections:
            ws.cell(row=current_row, column=1, value=sec_name)
            current_row += 1
            for r_vals in sec_rows:
                for col_idx, val in enumerate(r_vals, start=1):
                    ws.cell(row=current_row, column=col_idx, value=val)
                current_row += 1

    # Totals row
    if totals_row:
        for col_idx, val in enumerate(totals_row, start=1):
            ws.cell(row=current_row, column=col_idx, value=val)
        current_row += 1

    buf = io.BytesIO()
    wb.save(buf)
    wb.close()
    return buf.getvalue()


def sample_units():
    return [
        # Unit 101: 1BR, Occupied with Base Rent + Pet continuation + Parking continuation
        ["101", "1BR", 750, CANARIES[0], 1500.0, "RENT", 1450.0, 500.0, "01/01/2025", "12/31/2025", "", 0.0],
        ["", "", "", "", "", "PET", 35.0, "", "", "", "", ""],
        ["", "", "", "", "", "PARK", 75.0, "", "", "", "", ""],
        # Unit 102: 2BR, Notice with Base Rent + repeating unit ID continuation for Storage
        ["102", "2BR", 950, CANARIES[1], 1800.0, "BASE", 1800.0, 600.0, "06/01/2024", "05/31/2025", "05/31/2025", 50.0],
        ["102", "", "", "", "", "STOR", 50.0, "", "", "", "", ""],
        # Unit 103: Studio, Vacant
        ["103", "Studio", 500, "Vacant", 1200.0, "", "", "", "", "", "", 0.0],
    ]


def codes(result: dict) -> set[str]:
    return {issue["code"] for issue in result["issues"]}


def check_citations(obj, raw_bytes: bytes):
    expected_sha = hashlib.sha256(raw_bytes).hexdigest()
    if isinstance(obj, dict):
        if {"sheet", "row", "row_end", "column"}.issubset(obj.keys()):
            assert obj["source_sha256"] == expected_sha
            assert obj["sheet"] == 1
            assert 1 <= obj["row"] <= obj["row_end"]
            assert 1 <= obj["column"] <= 128
        for v in obj.values():
            check_citations(v, raw_bytes)
    elif isinstance(obj, list):
        for item in obj:
            check_citations(item, raw_bytes)


def test_positive_split_header_extraction_and_folding():
    raw = build_synthetic_workbook(unit_rows=sample_units())
    result = normalize_lease_charge_xlsx(raw)

    assert set(result.keys()) == ROOT_KEYS
    assert result["version"] == "1.0"
    assert result["pms_type"] == "realpage"
    assert result["source_sha256"] == hashlib.sha256(raw).hexdigest()

    # Floorplans "1BR", "2BR", "Studio" do NOT resolve use => explicit UNRESOLVED_UNIT_USE
    assert result["status"] == "blocked"
    assert result["residential_units"] == []
    assert result["commercial_units"] == []

    unresolved = [i["observation"] for i in result["issues"] if i["code"] == "UNRESOLVED_UNIT_USE"]
    # Physical unit count must be 3 (not inflated by charge continuations!)
    assert len(unresolved) == 3

    u101, u102, u103 = unresolved
    assert u101["unit_id"] == "101"
    assert u101["status"] == "occupied"
    assert u101["unit_type"] is None
    assert u101["tenant_name"] == "[REDACTED]"
    # Unit 101 has 3 evidence entries (primary row + 2 continuations: PET, PARK)
    assert len(u101["evidence"]) == 3
    assert "charge" in u101["evidence"][0]
    assert "charge" in u101["evidence"][1]
    assert "charge" in u101["evidence"][2]

    assert u102["unit_id"] == "102"
    assert u102["status"] == "occupied"
    # Unit 102 has 2 evidence entries (primary row + repeating unit continuation: STOR)
    assert len(u102["evidence"]) == 2

    assert u103["unit_id"] == "103"
    assert u103["status"] == "vacant"
    assert len(u103["evidence"]) == 1

    # Check issues
    assert "CONTINUATION_EVIDENCE" in codes(result)
    assert "UNRESOLVED_UNIT_USE" in codes(result)

    # Check exact citations
    check_citations(result, raw)

    # PII Check: No raw resident names leaked in serialized output
    encoded = json.dumps(result)
    for canary in CANARIES:
        assert canary not in encoded


def test_resolved_residential_and_commercial_use():
    unit_rows = [
        ["101", "Residential", 800, CANARIES[0], 1500.0, "RENT", 1500.0, 500.0, "01/01/2025", "12/31/2025", "", 0.0],
        ["102", "Residential", 850, "Vacant", 1550.0, "", "", "", "", "", "", 0.0],
        ["C1", "Commercial", 1200, CANARIES[1], 3000.0, "BASE", 3000.0, 1000.0, "01/01/2024", "12/31/2026", "", 0.0],
    ]
    raw = build_synthetic_workbook(unit_rows=unit_rows)
    result = normalize_lease_charge_xlsx(raw)

    assert result["status"] == "normalized_unvalidated"
    assert len(result["residential_units"]) == 2
    assert len(result["commercial_units"]) == 1

    assert result["counts"]["residential"] == {
        "occupied": 1,
        "vacant": 1,
        "down": 0,
        "total": 2,
    }
    assert result["counts"]["commercial"] == {
        "occupied": 1,
        "vacant": 0,
        "down": 0,
        "total": 1,
    }


def test_charge_code_classification():
    # Base rent codes
    for code in ("RENT", "Rent", "BASE", "Base Rent", "MRENT", "Market Rent", "Apt Rent", "APTR", "LEASE"):
        assert is_base_rent(code) is True
        assert is_ancillary_charge(code) is False
        assert classify_charge(code) == "base_rent"

    # Ancillary categories
    assert classify_charge("PARK") == "parking"
    assert classify_charge("Garage") == "parking"
    assert classify_charge("PET") == "pet"
    assert classify_charge("Pet Rent") == "pet"
    assert classify_charge("STOR") == "storage"
    assert classify_charge("Storage Locker") == "storage"
    assert classify_charge("UTIL") == "utilities"
    assert classify_charge("Electric") == "utilities"
    assert classify_charge("RUBS") == "utilities"
    assert classify_charge("W/D") == "washer_dryer"
    assert classify_charge("Washer") == "washer_dryer"
    assert classify_charge("CONC") == "concessions"
    assert classify_charge("Concession") == "concessions"

    for code in ("PARK", "PET", "STOR", "UTIL", "W/D", "CONC"):
        assert is_base_rent(code) is False
        assert is_ancillary_charge(code) is True


def test_blank_unit_and_repeating_continuations_do_not_inflate_count():
    unit_rows = [
        ["101", "1BR", 700, CANARIES[0], 1200.0, "RENT", 1200.0, 400.0, "01/01/2025", "12/31/2025", "", 0.0],
        ["", "", "", "", "", "PET", 25.0, "", "", "", "", ""],
        ["", "", "", "", "", "UTIL", 40.0, "", "", "", "", ""],
        ["101", "", "", "", "", "PARK", 50.0, "", "", "", "", ""],
        ["101", "", "", "", "", "W/D", 30.0, "", "", "", "", ""],
    ]
    raw = build_synthetic_workbook(unit_rows=unit_rows)
    result = normalize_lease_charge_xlsx(raw)

    unresolved = [i["observation"] for i in result["issues"] if i["code"] == "UNRESOLVED_UNIT_USE"]
    assert len(unresolved) == 1
    u = unresolved[0]
    assert u["unit_id"] == "101"
    # 1 primary row + 4 continuation rows = 5 evidence items
    assert len(u["evidence"]) == 5
    assert "CONTINUATION_EVIDENCE" in codes(result)


def test_future_and_applicants_section_excluded():
    current_units = [
        ["101", "1BR", 700, CANARIES[0], 1200.0, "RENT", 1200.0, 400.0, "01/01/2025", "12/31/2025", "", 0.0],
    ]
    extra_sections = [
        ("Future Residents", [
            ["102", "1BR", 750, "Future Tenant A", 1300.0, "RENT", 1300.0, 500.0, "03/01/2026", "02/28/2027", "", 0.0],
        ]),
        ("Applicants", [
            ["103", "2BR", 900, "Applicant B", 1600.0, "RENT", 1600.0, 600.0, "", "", "", 0.0],
        ]),
    ]
    raw = build_synthetic_workbook(unit_rows=current_units, extra_sections=extra_sections)
    result = normalize_lease_charge_xlsx(raw)

    unresolved = [i["observation"] for i in result["issues"] if i["code"] == "UNRESOLVED_UNIT_USE"]
    assert len(unresolved) == 1
    assert unresolved[0]["unit_id"] == "101"

    assert "NON_CURRENT_SECTION_EXCLUDED" in codes(result)
    assert "Future Tenant A" not in json.dumps(result)
    assert "Applicant B" not in json.dumps(result)


def test_totals_section_terminates_parsing():
    current_units = [
        ["101", "1BR", 700, CANARIES[0], 1200.0, "RENT", 1200.0, 400.0, "01/01/2025", "12/31/2025", "", 0.0],
    ]
    totals_row = ["Totals:", "", "", "", 1200.0, "", 1200.0, 400.0, "", "", "", 0.0]
    raw = build_synthetic_workbook(unit_rows=current_units, totals_row=totals_row)
    result = normalize_lease_charge_xlsx(raw)

    unresolved = [i["observation"] for i in result["issues"] if i["code"] == "UNRESOLVED_UNIT_USE"]
    assert len(unresolved) == 1
    assert unresolved[0]["unit_id"] == "101"


def test_status_mapping_and_down_units():
    unit_rows = [
        ["101", "Residential", 700, CANARIES[0], 1200.0, "RENT", 1200.0, 400.0, "01/01/2025", "12/31/2025", "", 0.0],
        ["102", "Residential", 750, "Notice", 1300.0, "RENT", 1300.0, 400.0, "01/01/2024", "12/31/2024", "12/31/2024", 0.0],
        ["103", "Residential", 800, "Vacant", 1400.0, "", "", "", "", "", "", 0.0],
        ["104", "Residential", 850, "Down", 0.0, "", "", "", "", "", "", 0.0],
    ]
    raw = build_synthetic_workbook(unit_rows=unit_rows)
    result = normalize_lease_charge_xlsx(raw)

    assert result["counts"]["residential"] == {
        "occupied": 2,  # 101 (current), 102 (notice)
        "vacant": 1,    # 103
        "down": 1,      # 104
        "total": 4,
    }


def test_orphan_continuation_emits_blocker():
    # Continuation row appears before any physical unit row
    unit_rows = [
        ["", "", "", "", "", "PET", 35.0, "", "", "", "", ""],
        ["101", "1BR", 700, CANARIES[0], 1200.0, "RENT", 1200.0, 400.0, "01/01/2025", "12/31/2025", "", 0.0],
    ]
    raw = build_synthetic_workbook(unit_rows=unit_rows)
    result = normalize_lease_charge_xlsx(raw)

    assert "ORPHAN_CONTINUATION" in codes(result)
    assert result["status"] == "blocked"


def test_duplicate_unit_conflict_blocks():
    # Same unit ID 101 with conflicting status
    unit_rows = [
        ["101", "1BR", 700, CANARIES[0], 1200.0, "RENT", 1200.0, 400.0, "01/01/2025", "12/31/2025", "", 0.0],
        ["101", "1BR", 700, "Vacant", 1200.0, "", "", "", "", "", "", 0.0],
    ]
    raw = build_synthetic_workbook(unit_rows=unit_rows)
    result = normalize_lease_charge_xlsx(raw)

    assert "DUPLICATE_UNIT_CONFLICT" in codes(result)
    assert result["status"] == "blocked"


def test_split_header_folding_alternate_distributions():
    # Variation where Row 5 has separate "Resident" and "Name" columns
    # and Row 6 has Deposit under Resident
    row5 = [
        "Unit", "Unit Type", "Unit", "Resident", "Name", "Market",
        "Charge", "Amount", "Resident", "Move In", "Lease",
        "Move Out", "Balance",
    ]
    row6 = [
        "", "", "Sq Ft", "", "", "Rent",
        "Code", "", "Deposit", "", "Expiration",
        "", "",
    ]
    unit_rows = [
        ["101", "1BR", 750, "Smith", "Alice", 1400.0, "RENT", 1400.0, 500.0, "01/01/2025", "12/31/2025", "", 0.0],
    ]
    raw = build_synthetic_workbook(row5_headers=row5, row6_headers=row6, unit_rows=unit_rows)
    result = normalize_lease_charge_xlsx(raw)

    unresolved = [i["observation"] for i in result["issues"] if i["code"] == "UNRESOLVED_UNIT_USE"]
    assert len(unresolved) == 1
    assert unresolved[0]["unit_id"] == "101"
    assert unresolved[0]["tenant_name"] == "[REDACTED]"


@pytest.mark.parametrize("payload", [b"", b"Not a zip file", b"PK\x03\x04corrupt_zip_content", "string_not_bytes", None])
def test_requires_valid_ooxml_bytes(payload):
    with pytest.raises(RentRollNormalizationError) as exc_info:
        normalize_lease_charge_xlsx(payload)
    assert exc_info.value.code in ("UNSUPPORTED_XLSX_FORMAT", "MALFORMED_INPUT", "EMPTY_INPUT")
    assert exc_info.value.__context__ is None
    assert exc_info.value.__cause__ is None


def test_missing_sheet_title_fails_layout():
    raw = build_synthetic_workbook(title="Other Sheet Title", unit_rows=sample_units())
    with pytest.raises(RentRollNormalizationError) as exc_info:
        normalize_lease_charge_xlsx(raw)
    assert exc_info.value.code == "UNSUPPORTED_XLSX_LAYOUT"


def test_missing_split_headers_fails_layout():
    bad_row5 = ["Col A", "Col B", "Col C", "Col D"]
    bad_row6 = ["", "", "", ""]
    raw = build_synthetic_workbook(row5_headers=bad_row5, row6_headers=bad_row6, unit_rows=sample_units())
    with pytest.raises(RentRollNormalizationError) as exc_info:
        normalize_lease_charge_xlsx(raw)
    assert exc_info.value.code == "UNSUPPORTED_XLSX_LAYOUT"


def test_missing_section_header_fails_layout():
    raw = build_synthetic_workbook(section_header="Invalid Section Header", unit_rows=sample_units())
    with pytest.raises(RentRollNormalizationError) as exc_info:
        normalize_lease_charge_xlsx(raw)
    assert exc_info.value.code == "UNSUPPORTED_XLSX_LAYOUT"


def test_vba_macro_rejected():
    # Build a zip file with a vbaproject entry
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("xl/vbaProject.bin", b"fake macro content")
        zf.writestr("[Content_Types].xml", b"<xml></xml>")
    raw = buf.getvalue()

    with pytest.raises(RentRollNormalizationError) as exc_info:
        normalize_lease_charge_xlsx(raw)
    assert exc_info.value.code == "MALFORMED_INPUT"


def test_input_byte_limit_exceeded():
    huge_bytes = b"PK\x03\x04" + b"\x00" * (33 * 1024 * 1024)
    with pytest.raises(RentRollNormalizationError) as exc_info:
        normalize_lease_charge_xlsx(huge_bytes)
    assert exc_info.value.code == "INPUT_LIMIT_EXCEEDED"


def test_citation_coordinate_accuracy():
    unit_rows = [
        ["201", "1BR", 650, CANARIES[0], 1100.0, "RENT", 1100.0, 350.0, "02/01/2025", "01/31/2026", "", 0.0],
        ["", "", "", "", "", "PET", 30.0, "", "", "", "", ""],
    ]
    raw = build_synthetic_workbook(unit_rows=unit_rows)
    result = normalize_lease_charge_xlsx(raw)

    unresolved = [i["observation"] for i in result["issues"] if i["code"] == "UNRESOLVED_UNIT_USE"]
    u = unresolved[0]
    # Unit 201 is at row 8, column 1
    assert u["evidence"][0]["unit_id"]["row"] == 8
    assert u["evidence"][0]["unit_id"]["column"] == 1
    # Unit type is at row 8, column 2
    assert u["evidence"][0]["unit_type"]["row"] == 8
    assert u["evidence"][0]["unit_type"]["column"] == 2
    # Base rent charge is at row 8, column 6
    assert u["evidence"][0]["charge"]["row"] == 8
    assert u["evidence"][0]["charge"]["column"] == 6
    # Continuation pet charge is at row 9, column 6
    assert u["evidence"][1]["charge"]["row"] == 9
    assert u["evidence"][1]["charge"]["column"] == 6


def test_multi_sheet_workbook_ignores_non_inventory_sheet():
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "Cover Sheet"
    ws1.cell(row=1, column=1, value="Property Overview")
    ws1.cell(row=2, column=1, value="Confidential")

    ws2 = wb.create_sheet(title="Rent Roll Detail")
    ws2.cell(row=1, column=1, value="Rent Roll with Lease Charges")
    row5 = ["Unit", "Unit Type", "Unit", "Resident Name", "Market", "Charge", "Amount", "Resident", "Move In", "Lease", "Move Out", "Balance"]
    row6 = ["", "", "Sq Ft", "", "Rent", "Code", "", "Deposit", "", "Expiration", "", ""]
    for c, v in enumerate(row5, 1):
        ws2.cell(row=5, column=c, value=v)
    for c, v in enumerate(row6, 1):
        ws2.cell(row=6, column=c, value=v)
    ws2.cell(row=7, column=1, value="Current/Notice/Vacant Residents")
    unit_row = ["101", "Residential", 750, CANARIES[0], 1400.0, "RENT", 1400.0, 500.0, "01/01/2025", "12/31/2025", "", 0.0]
    for c, v in enumerate(unit_row, 1):
        ws2.cell(row=8, column=c, value=v)

    buf = io.BytesIO()
    wb.save(buf)
    wb.close()
    raw = buf.getvalue()

    result = normalize_lease_charge_xlsx(raw)
    assert "NON_INVENTORY_SHEET_IGNORED" in codes(result)
    assert len(result["residential_units"]) == 1
    # Sheet index cited must be 2
    assert result["residential_units"][0]["evidence"][0]["unit_id"]["sheet"] == 2


def test_v2_compat_projection():
    from plat_harness.ingest.compat import _project
    from plat_harness.ingest.contracts import validate_observations

    raw = build_synthetic_workbook(unit_rows=sample_units())
    v1 = normalize_lease_charge_xlsx(raw)
    source_id = "src_" + "a" * 32
    v2 = _project(v1, source_id, None, None, "onesite-detailed-realpage")

    # Contract validation succeeds
    assert validate_observations(v2, subject_id=None, as_of=None) == v2
    assert v2["status"] == "blocked"
    assert len(v2["unknown_use_units"]) == 3
