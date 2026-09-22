"""Synthetic regressions for the flat normalizer's fail-closed boundaries.

No real exports, metadata, resident values, or source paths are fixtures.
"""
from __future__ import annotations

import csv
import hashlib
import io
import json
import traceback
import warnings
import zipfile

import pytest

from plat_harness.ingest import pms_normalizer as normalizer

CANARY = "SYNTHETIC-PRIVATE-CONTEXT-CANARY"
HEADERS = {
    "yardi": ["Unit", "Status", "Unit Type"],
    "realpage": ["Unit #", "Occupancy Status", "Unit Category"],
    "entrata": ["Unit Number", "Occupancy", "Space Type"],
}
NULLS = dict.fromkeys(("occupied", "vacant", "down", "total"))


def csv_input(rows):
    stream = io.StringIO(newline="")
    csv.writer(stream).writerows(rows)
    return io.StringIO(stream.getvalue())


def workbook_input(tables, *, number_format=None):
    from openpyxl import Workbook
    book = Workbook()
    book.properties.creator = CANARY
    for index, rows in enumerate(tables):
        sheet = book.active if index == 0 else book.create_sheet()
        sheet.title = "Synthetic" + str(index)
        for row in rows:
            sheet.append(row)
        if number_format is not None and index == 0:
            sheet["A2"].number_format = number_format
    stream = io.BytesIO()
    book.save(stream)
    book.close()
    stream.seek(0)
    return stream


def assert_private(result, capsys, caplog):
    captured = capsys.readouterr()
    assert CANARY not in json.dumps(result)
    assert CANARY not in captured.out + captured.err + caplog.text
    for issue in result["issues"]:
        assert set(issue) == {"code", "severity", "citation"}
    for scope in ("residential", "commercial"):
        for unit in result[scope + "_units"]:
            assert set(unit) <= {"unit_id", "status", "unit_type", "tenant_name", "phone", "email", "evidence"}
            for evidence in unit["evidence"]:
                assert not {"property", "building"}.intersection(evidence)


def assert_blocked(result, code):
    assert result["status"] == "blocked"
    assert result["counts"] == {"residential": NULLS, "commercial": NULLS}
    assert code in {item["code"] for item in result["issues"]}


@pytest.mark.parametrize("pms", HEADERS)
@pytest.mark.parametrize("field", ["Building", "Property"])
@pytest.mark.parametrize("format_", ["csv", "xlsx"])
def test_explicit_multiple_contexts_do_not_collapse_identity(pms, field, format_, capsys, caplog):
    rows = [HEADERS[pms] + [field, "Resident"],
            ["101", "Current", "Residential", CANARY + "-A", CANARY],
            ["101", "Current", "Residential", CANARY + "-B", CANARY]]
    stream = csv_input(rows) if format_ == "csv" else workbook_input([rows])
    result = normalizer.normalize_rent_roll(stream, pms)
    assert_private(result, capsys, caplog)
    assert_blocked(result, "AMBIGUOUS_UNIT_CONTEXT")
    assert result["residential_units"] == result["commercial_units"] == []


@pytest.mark.parametrize("case", ["blank", "ragged", "header_removed", "header_added", "formula"])
def test_missing_or_changed_context_cannot_prove_duplicate_identity(case, capsys, caplog):
    base = HEADERS["yardi"]
    rows = [base + ["Building"], ["101", "Current", "Residential", CANARY]]
    if case == "header_removed":
        rows += [base, ["101", "Current", "Residential"]]
    elif case == "header_added":
        rows = [base, ["101", "Current", "Residential"], *rows]
    else:
        row = ["101", "Current", "Residential"]
        if case != "ragged":
            row.append("=" + CANARY if case == "formula" else "")
        rows.append(row)
    result = normalizer.normalize_rent_roll(csv_input(rows), "yardi")
    assert_private(result, capsys, caplog)
    assert result["status"] == "blocked"
    assert result["counts"] == {"residential": NULLS, "commercial": NULLS}
    assert result["residential_units"] == []


@pytest.mark.parametrize("field", ["Building", "Building ID", "Building Number", "Building #", "Building Name", "Property", "Property ID", "Property Code", "Property Name"])
def test_context_aliases_single_scope_dedupe_and_redaction(field, capsys, caplog):
    rows = [HEADERS["yardi"] + [field, "Resident", "Record Type"],
            ["101", "Current", "Residential", CANARY, CANARY, "unit"],
            ["101", "Current", "Residential", CANARY, CANARY, "unit"],
            ["101", "", "", CANARY, CANARY, "charge"]]
    result = normalizer.normalize_rent_roll(csv_input(rows), "yardi")
    assert_private(result, capsys, caplog)
    assert result["status"] == "normalized_unvalidated"
    assert result["counts"]["residential"]["total"] == 1
    unit = result["residential_units"][0]
    assert unit["tenant_name"] == "[REDACTED]"
    assert len(unit["evidence"]) == 3
    # A second context must be recognized for every supported alias.
    rows.append(["102", "Vacant", "Residential", CANARY + "-B", CANARY, "unit"])
    blocked = normalizer.normalize_rent_roll(csv_input(rows), "yardi")
    assert_private(blocked, capsys, caplog)
    assert_blocked(blocked, "AMBIGUOUS_UNIT_CONTEXT")


def test_property_and_building_context_checked_across_sheets(capsys, caplog):
    header = HEADERS["yardi"] + ["Property", "Building"]
    first = [header, ["101", "Current", "Residential", CANARY, "A"]]
    second = [header, ["102", "Vacant", "Residential", CANARY, "B"]]
    result = normalizer.normalize_rent_roll(workbook_input([first, second]), "yardi")
    assert_private(result, capsys, caplog)
    assert_blocked(result, "AMBIGUOUS_UNIT_CONTEXT")


@pytest.mark.parametrize("number_format", ["0000", "0", "0.00", "0.00E+00", "@", '"' + CANARY + '"0000'])
def test_numeric_formatted_identifier_rejected_not_silently_rewritten(number_format, capsys, caplog):
    rows = [HEADERS["yardi"], [101, "Current", "Residential"], ["0101", "Vacant", "Residential"]]
    stream = workbook_input([rows], number_format=number_format)
    digest = hashlib.sha256(stream.getvalue()).hexdigest()
    result = normalizer.normalize_rent_roll(stream, "yardi")
    assert_private(result, capsys, caplog)
    assert_blocked(result, "UNSUPPORTED_UNIT_ID_FORMAT")
    assert {unit["unit_id"] for unit in result["residential_units"]} == {"0101"}
    issue = next(item for item in result["issues"] if item["code"] == "UNSUPPORTED_UNIT_ID_FORMAT")
    assert issue["citation"] == {"source_sha256": digest, "sheet": 1, "row": 2, "row_end": 2, "column": 1}


@pytest.mark.parametrize("identifier,format_", [(101, "General"), ("0101", "0000")])
def test_general_numeric_and_text_identifiers_retain_physical_citation(identifier, format_):
    stream = workbook_input([[HEADERS["yardi"], [identifier, "Current", "Residential"]]], number_format=format_)
    result = normalizer.normalize_rent_roll(stream, "yardi")
    assert result["status"] == "normalized_unvalidated"
    unit = result["residential_units"][0]
    assert unit["unit_id"] == str(identifier)
    assert unit["evidence"][0]["unit_id"]["row"] == 2


def test_numeric_format_in_discarded_amount_column_does_not_block():
    from openpyxl import load_workbook
    stream = workbook_input([[HEADERS["yardi"] + ["Rent"], ["0101", "Current", "Residential", 1234]]])
    book = load_workbook(stream)
    book.active["D2"].number_format = "0.00"
    updated = io.BytesIO()
    book.save(updated)
    book.close()
    updated.seek(0)
    assert normalizer.normalize_rent_roll(updated, "yardi")["status"] == "normalized_unvalidated"


@pytest.mark.parametrize("pms", HEADERS)
@pytest.mark.parametrize("missing", [0, 1, 2])
@pytest.mark.parametrize("position", ["before", "after"])
def test_partial_inventory_sheet_blocks_complete_table(pms, missing, position, capsys, caplog):
    headers = HEADERS[pms]
    complete = [headers, ["101", "Current", "Residential"]]
    partial = [[value for i, value in enumerate(headers) if i != missing] + ["Resident"],
               [value for i, value in enumerate(["102", "Vacant", "Residential"]) if i != missing] + [CANARY]]
    tables = [partial, complete] if position == "before" else [complete, partial]
    result = normalizer.normalize_rent_roll(workbook_input(tables), pms)
    assert_private(result, capsys, caplog)
    assert_blocked(result, "INCOMPLETE_INVENTORY_HEADER")


def test_partial_repeated_csv_header_blocks_and_cannot_heal():
    rows = [HEADERS["yardi"], ["101", "Current", "Residential"],
            ["Unit", "Status"], ["102", "Vacant"],
            HEADERS["yardi"], ["103", "Down", "Residential"],
            ["Scope", "Occupied", "Vacant", "Down", "Total"], ["Residential", 1, 0, 1, 2]]
    result = normalizer.normalize_rent_roll(csv_input(rows), "yardi")
    assert_blocked(result, "INCOMPLETE_INVENTORY_HEADER")
    assert result["summary"]["residential"]["status"] == "unresolved"


def test_arbitrary_metadata_sheet_still_ignored(capsys, caplog):
    tables = [[["Prepared By", "Date"], [CANARY, CANARY]],
              [HEADERS["yardi"], ["101", "Current", "Residential"]]]
    result = normalizer.normalize_rent_roll(workbook_input(tables), "yardi")
    assert_private(result, capsys, caplog)
    assert result["status"] == "normalized_unvalidated"
    assert "NON_INVENTORY_SHEET_IGNORED" in {i["code"] for i in result["issues"]}


@pytest.mark.parametrize("partial", [
    [["Unit", "Rent"], ["102", "1000"]],
    [["Rent", "Unit"], ["1000", "102"]],
])
def test_unit_header_with_unknown_columns_is_not_metadata(partial):
    complete = [HEADERS["yardi"], ["101", "Current", "Residential"]]
    result = normalizer.normalize_rent_roll(workbook_input([complete, partial]), "yardi")
    assert_blocked(result, "INCOMPLETE_INVENTORY_HEADER")


@pytest.mark.parametrize("limit", ["_MAX_ROWS", "_MAX_CELLS"])
def test_workbook_cumulative_bounds_checked_before_first_iterator(limit, monkeypatch):
    from openpyxl.worksheet._read_only import ReadOnlyWorksheet
    rows = [HEADERS["yardi"], ["101", "Current", "Residential"]]
    stream = workbook_input([rows, rows])
    monkeypatch.setattr(normalizer, limit, 3 if limit == "_MAX_ROWS" else 10)
    calls = []
    def forbidden_iterator(*args, **kwargs):
        calls.append(True)
        return iter(())
    monkeypatch.setattr(ReadOnlyWorksheet, "iter_rows", forbidden_iterator)
    with pytest.raises(normalizer.RentRollNormalizationError) as error:
        normalizer.normalize_rent_roll(stream, "yardi")
    assert error.value.code == "INPUT_LIMIT_EXCEEDED"
    assert calls == []


def test_duplicate_context_header_rejected_without_value_echo():
    rows = [HEADERS["yardi"] + ["Property", "Property ID"],
            ["101", "Current", "Residential", CANARY, CANARY]]
    with pytest.raises(normalizer.RentRollNormalizationError) as error:
        normalizer.normalize_rent_roll(csv_input(rows), "yardi")
    assert error.value.code == "AMBIGUOUS_HEADER"
    assert CANARY not in str(error.value)
    assert error.value.__context__ is None


def replaced_sheet(body):
    original = workbook_input([[HEADERS["yardi"], ["101", "Current", "Residential"]]])
    replacement = ('<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
                   '<dimension ref="A1:A1"/><sheetData>' + body + '</sheetData></worksheet>').encode()
    stream = io.BytesIO()
    with zipfile.ZipFile(original) as source, zipfile.ZipFile(stream, "w") as target:
        for item in source.infolist():
            target.writestr(item, replacement if item.filename == "xl/worksheets/sheet1.xml" else source.read(item.filename))
    stream.seek(0)
    return stream


@pytest.mark.parametrize("body", [
    '<row r="1"><c r="A1"/><c r="A1"/></row>',
    '<row r="1"><c r="B1"/><c r="A1"/></row>',
    '<row r="1"><c r="A2"/></row>',
    '<row r="1"/><row r="1"/>',
    '<row r="1"><other r="XFD1"/></row>',
])
def test_malformed_coordinate_order_does_not_overwrite_or_relocate(body):
    with pytest.raises(normalizer.RentRollNormalizationError) as error:
        normalizer.normalize_rent_roll(replaced_sheet(body), "yardi")
    assert error.value.code == "MALFORMED_INPUT"
    assert error.value.__context__ is None


def test_small_real_sparse_row_hits_limit_without_iterator_mock(monkeypatch):
    stream = replaced_sheet('<row r="1"><c r="E1"><v>1</v></c></row>')
    monkeypatch.setattr(normalizer, "_MAX_COLUMNS", 4)
    with pytest.raises(normalizer.RentRollNormalizationError) as error:
        normalizer.normalize_rent_roll(stream, "yardi")
    assert error.value.code == "INPUT_LIMIT_EXCEEDED"


@pytest.mark.parametrize("body", [
    '<row r="1"><c r="XFD1"><v>1</v></c></row>',
    '<row r="50001"><c r="A50001"><v>1</v></c></row>',
    '<row r="1"><c r="A50001"><v>1</v></c></row>',
    '<row r="1">' + '<c><v>1</v></c>' * 129 + '</row>',
])
def test_actual_xml_bounds_rejected_before_any_row_iterator(body, monkeypatch, capsys, caplog):
    from openpyxl.worksheet._read_only import ReadOnlyWorksheet
    stream = replaced_sheet(body)
    calls = []
    def forbidden_iterator(*args, **kwargs):
        calls.append(True)
        return iter(())  # Never allocate hostile rows, even on the vulnerable baseline.
    monkeypatch.setattr(ReadOnlyWorksheet, "iter_rows", forbidden_iterator)
    with warnings.catch_warnings(record=True) as emitted:
        warnings.simplefilter("always")
        with pytest.raises(normalizer.RentRollNormalizationError) as error:
            normalizer.normalize_rent_roll(stream, "yardi")
    assert emitted == []
    captured = capsys.readouterr()
    assert CANARY not in captured.out + captured.err + caplog.text + "".join(traceback.format_exception(error.value))
    assert error.value.__context__ is None
    assert error.value.code == "INPUT_LIMIT_EXCEEDED"
    assert calls == []


@pytest.mark.parametrize("limit,body", [
    ("_MAX_ROWS", '<row/><row/><row/><row/><row/>'),
    ("_MAX_CELLS", '<row r="1"><c r="D1"><v>1</v></c></row><row r="2"><c r="D2"><v>1</v></c></row>'),
])
def test_inferred_rows_and_sparse_padded_cell_budget_preflight(limit, body, monkeypatch):
    from openpyxl.worksheet._read_only import ReadOnlyWorksheet
    stream = replaced_sheet(body)
    monkeypatch.setattr(normalizer, limit, 4)
    calls = []
    def forbidden_iterator(*args, **kwargs):
        calls.append(True)
        return iter(())
    monkeypatch.setattr(ReadOnlyWorksheet, "iter_rows", forbidden_iterator)
    with pytest.raises(normalizer.RentRollNormalizationError) as error:
        normalizer.normalize_rent_roll(stream, "yardi")
    assert error.value.code == "INPUT_LIMIT_EXCEEDED"
    assert calls == []
