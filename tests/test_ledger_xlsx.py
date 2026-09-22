"""Synthetic-only contract tests for unknown-vendor ledger/charge-detail XLSX.

All workbooks are generated in-memory via openpyxl. No live or private deal files.
Sheet names are arbitrary; inventory is selected by headers, never by title.
"""
from __future__ import annotations

import hashlib
import io
import json
import traceback
import zipfile

import pytest
from openpyxl import Workbook

from plat_harness.ingest.ledger_xlsx import (
    ANCILLARY_CATEGORIES,
    ANCILLARY_CODES,
    BASE_RENT_CODES,
    classify_charge,
    is_ancillary_charge,
    is_base_rent,
    normalize_ledger_xlsx,
)
from plat_harness.ingest.pms_normalizer import RentRollNormalizationError

CANARIES = (
    "SYNTHETIC CANARY TENANT",
    "John Doe Canary",
    "Jane Smith Canary",
    "canary_user@example.invalid",
    "+1-555-019-9999",
    "SECRET_PII_NOTE",
)

ROOT_KEYS = {
    "version",
    "pms_type",
    "source_sha256",
    "residential_units",
    "commercial_units",
    "counts",
    "summary",
    "issues",
    "status",
}

DEFAULT_HEADERS = [
    "Bldg-Unit",
    "Unit Type",
    "SQFT",
    "Unit Status",
    "Resident",
    "Market Rent",
    "Ledger",
    "Charge Code",
    "Scheduled Charges",
    "Balance",
    "Deposit Held",
    "Move-In",
    "Lease Start",
    "Lease End",
    "Expected Move-Out",
]


def build_synthetic_workbook(
    *,
    headers: list[str] | None = None,
    header_row: int = 7,
    unit_rows: list[list] | None = None,
    extra_sections: list[tuple[str, list[list]]] | None = None,
    totals_row: list | None = None,
    inventory_sheet_title: str = "InventoryDetail",
    parameters_first: bool = True,
    extra_columns: int = 0,
    include_parameters: bool = True,
) -> bytes:
    wb = Workbook()
    if include_parameters and parameters_first:
        params = wb.active
        params.title = "Report Parameters"
        inv = wb.create_sheet(inventory_sheet_title)
    elif include_parameters:
        inv = wb.active
        inv.title = inventory_sheet_title
        params = wb.create_sheet("Report Parameters")
    else:
        params = None
        inv = wb.active
        inv.title = inventory_sheet_title
    if params is not None:
        params.cell(row=1, column=1, value="Parameter")
        params.cell(row=1, column=2, value="Value")
        params.cell(row=2, column=1, value="As of")
        params.cell(row=2, column=2, value="01/15/2026")
        params.cell(row=3, column=1, value="Include Zero Dollar Charges")
        params.cell(row=3, column=2, value="Yes")

    inv.cell(row=1, column=1, value="Ledger / Charge Detail")
    inv.cell(row=2, column=1, value="Property: Synthetic Mixed-Use Asset")
    inv.cell(row=3, column=1, value="As of Date: 01/15/2026")

    hdr = list(headers if headers is not None else DEFAULT_HEADERS)
    for col_idx, val in enumerate(hdr, start=1):
        inv.cell(row=header_row, column=col_idx, value=val)
    for extra in range(extra_columns):
        inv.cell(row=header_row, column=len(hdr) + extra + 1, value=f"Extra-{extra}")

    current_row = header_row + 1
    if unit_rows is not None:
        for r_vals in unit_rows:
            for col_idx, val in enumerate(r_vals, start=1):
                inv.cell(row=current_row, column=col_idx, value=val)
            current_row += 1

    if extra_sections:
        for sec_name, sec_rows in extra_sections:
            inv.cell(row=current_row, column=1, value=sec_name)
            current_row += 1
            for r_vals in sec_rows:
                for col_idx, val in enumerate(r_vals, start=1):
                    inv.cell(row=current_row, column=col_idx, value=val)
                current_row += 1

    if totals_row:
        for col_idx, val in enumerate(totals_row, start=1):
            inv.cell(row=current_row, column=col_idx, value=val)

    buf = io.BytesIO()
    wb.save(buf)
    wb.close()
    return buf.getvalue()


def sample_units():
    return [
        ["A-101", "1BR", 750, "Occupied", CANARIES[0], 1500.0, "Resident", "RENT", 1450.0, 0.0, 500.0, "01/01/2025", "01/01/2025", "12/31/2025", ""],
        ["", "", "", "", "", "", "Resident", "PET", 35.0, "", "", "", "", "", ""],
        ["A-101", "", "", "", "", "", "Subsidy", "PARK", 75.0, "", "", "", "", "", ""],
        ["B-101", "1BR", 750, "Occupied", CANARIES[1], 1500.0, "Resident", "BASE", 1500.0, 0.0, 500.0, "02/01/2025", "02/01/2025", "01/31/2026", ""],
        ["A-201", "Studio", 500, "Vacant", "", 1200.0, "", "", "", 0.0, 0.0, "", "", "", ""],
        ["C-10", "Retail", 1100, "Occupied", CANARIES[2], 3200.0, "Resident", "RENT", 3200.0, 0.0, 1000.0, "01/01/2024", "01/01/2024", "12/31/2026", ""],
        ["C-10", "", "", "", "", "", "CAM", "UTIL", 200.0, "", "", "", "", "", ""],
    ]


def codes(result: dict) -> set[str]:
    return {issue["code"] for issue in result["issues"]}


def unresolved_units(result: dict) -> list[dict]:
    return [i["observation"] for i in result["issues"] if i["code"] == "UNRESOLVED_UNIT_USE"]


def check_citations(obj, raw_bytes: bytes) -> None:
    expected_sha = hashlib.sha256(raw_bytes).hexdigest()
    if isinstance(obj, dict):
        if {"sheet", "row", "row_end", "column"}.issubset(obj.keys()):
            assert obj["source_sha256"] == expected_sha
            assert isinstance(obj["sheet"], int) and obj["sheet"] >= 1
            assert 1 <= obj["row"] <= obj["row_end"]
            assert 1 <= obj["column"] <= 128
            dumped = json.dumps(obj)
            for canary in CANARIES:
                assert canary not in dumped
        for v in obj.values():
            check_citations(v, raw_bytes)
    elif isinstance(obj, list):
        for item in obj:
            check_citations(item, raw_bytes)


def assert_no_pii(result: dict) -> None:
    encoded = json.dumps(result)
    for canary in CANARIES:
        assert canary not in encoded
    assert "Advenir" not in encoded
    assert "Station 121" not in encoded


def test_positive_header_selection_not_sheet_name():
    raw = build_synthetic_workbook(
        inventory_sheet_title="NotAPropertyName",
        parameters_first=True,
        unit_rows=sample_units(),
    )
    result = normalize_ledger_xlsx(raw)

    assert set(result.keys()) == ROOT_KEYS
    assert result["version"] == "1.0"
    assert result["pms_type"] == "unknown"
    assert result["source_sha256"] == hashlib.sha256(raw).hexdigest()
    assert result["status"] == "blocked"
    assert "NON_INVENTORY_SHEET_IGNORED" in codes(result)

    unresolved = unresolved_units(result)
    # Four physical identities: A-101, B-101, A-201, C-10 (C-10 is Retail -> commercial, not unresolved)
    assert [u["unit_id"] for u in unresolved] == ["A-101", "B-101", "A-201"]
    assert len(result["commercial_units"]) == 1
    assert result["commercial_units"][0]["unit_id"] == "C-10"
    assert result["commercial_units"][0]["unit_type"] == "commercial"
    assert result["commercial_units"][0]["tenant_name"] == "[REDACTED]"
    assert result["residential_units"] == []
    assert all(v is None for v in result["counts"]["residential"].values())
    assert all(v is None for v in result["counts"]["commercial"].values())

    u101 = unresolved[0]
    assert u101["unit_id"] == "A-101"
    assert u101["status"] == "occupied"
    assert u101["unit_type"] is None
    assert u101["tenant_name"] == "[REDACTED]"
    # primary + blank continuation (PET) + repeating identity ledger (PARK)
    assert len(u101["evidence"]) == 3
    assert "CONTINUATION_EVIDENCE" in codes(result)

    b101 = unresolved[1]
    assert b101["unit_id"] == "B-101"
    assert b101["status"] == "occupied"

    assert unresolved[2]["unit_id"] == "A-201"
    assert unresolved[2]["status"] == "vacant"

    check_citations(result, raw)
    assert_no_pii(result)


def test_report_parameters_never_become_inventory():
    raw = build_synthetic_workbook(
        inventory_sheet_title="ChargeLedger",
        parameters_first=True,
        unit_rows=[["A-101", "Residential", 700, "Occupied", CANARIES[0], 1200.0, "Resident", "RENT", 1200.0, 0.0, 400.0, "01/01/2025", "01/01/2025", "12/31/2025", ""]],
    )
    result = normalize_ledger_xlsx(raw)
    assert result["status"] == "normalized_unvalidated"
    assert len(result["residential_units"]) == 1
    assert result["residential_units"][0]["unit_id"] == "A-101"
    assert result["commercial_units"] == []
    assert "NON_INVENTORY_SHEET_IGNORED" in codes(result)
    assert result["residential_units"][0]["evidence"][0]["unit_id"]["sheet"] == 2
    assert_no_pii(result)


def test_inventory_found_when_parameters_sheet_is_second():
    raw = build_synthetic_workbook(
        inventory_sheet_title="Units",
        parameters_first=False,
        unit_rows=[["B-9", "Apartment", 640, "Current", CANARIES[0], 1100.0, "Resident", "RENT", 1100.0, 0.0, 300.0, "01/01/2025", "01/01/2025", "12/31/2025", ""]],
    )
    result = normalize_ledger_xlsx(raw)
    assert len(result["residential_units"]) == 1
    assert result["residential_units"][0]["unit_id"] == "B-9"
    assert result["residential_units"][0]["evidence"][0]["unit_id"]["sheet"] == 1
    assert "NON_INVENTORY_SHEET_IGNORED" in codes(result)


def test_header_row_is_discovered_not_hardcoded():
    raw = build_synthetic_workbook(
        header_row=5,
        inventory_sheet_title="Detail",
        unit_rows=[["A-5", "Residential", 800, "Occupied", CANARIES[0], 1400.0, "Resident", "RENT", 1400.0, 0.0, 400.0, "01/01/2025", "01/01/2025", "12/31/2025", ""]],
    )
    result = normalize_ledger_xlsx(raw)
    assert result["residential_units"][0]["evidence"][0]["unit_id"]["row"] == 6
    assert result["residential_units"][0]["unit_id"] == "A-5"


def test_building_relative_ids_remain_distinct():
    unit_rows = [
        ["A-101", "Residential", 700, "Occupied", CANARIES[0], 1200.0, "Resident", "RENT", 1200.0, 0.0, 400.0, "01/01/2025", "01/01/2025", "12/31/2025", ""],
        ["B-101", "Residential", 700, "Vacant", "", 1200.0, "", "", "", 0.0, 0.0, "", "", "", ""],
    ]
    raw = build_synthetic_workbook(unit_rows=unit_rows)
    result = normalize_ledger_xlsx(raw)
    ids = [u["unit_id"] for u in result["residential_units"]]
    assert ids == ["A-101", "B-101"]
    assert result["counts"]["residential"] == {
        "occupied": 1,
        "vacant": 1,
        "down": 0,
        "total": 2,
    }


def test_continuations_and_ledgers_do_not_inflate_unit_count():
    unit_rows = [
        ["A-101", "1BR", 700, "Occupied", CANARIES[0], 1200.0, "Resident", "RENT", 1200.0, 0.0, 400.0, "01/01/2025", "01/01/2025", "12/31/2025", ""],
        ["", "", "", "", "", "", "Resident", "PET", 25.0, "", "", "", "", "", ""],
        ["", "", "", "", "", "", "Resident", "UTIL", 40.0, "", "", "", "", "", ""],
        ["A-101", "", "", "", "", "", "Other", "PARK", 50.0, "", "", "", "", "", ""],
        ["A-101", "1BR", 700, "Occupied", CANARIES[0], 1200.0, "Subsidy", "STOR", 30.0, "", "", "", "", "", ""],
    ]
    raw = build_synthetic_workbook(unit_rows=unit_rows)
    result = normalize_ledger_xlsx(raw)
    unresolved = unresolved_units(result)
    assert len(unresolved) == 1
    assert unresolved[0]["unit_id"] == "A-101"
    assert len(unresolved[0]["evidence"]) == 5
    assert "CONTINUATION_EVIDENCE" in codes(result)


def test_labelled_retail_cannot_disappear():
    unit_rows = [
        ["A-1", "Residential", 700, "Occupied", CANARIES[0], 1100.0, "Resident", "RENT", 1100.0, 0.0, 300.0, "01/01/2025", "01/01/2025", "12/31/2025", ""],
        ["R-1", "Retail", 900, "Occupied", CANARIES[1], 2500.0, "Resident", "RENT", 2500.0, 0.0, 800.0, "01/01/2024", "01/01/2024", "12/31/2026", ""],
        ["R-2", "Commercial", 800, "Vacant", "", 2400.0, "", "", "", 0.0, 0.0, "", "", "", ""],
        ["O-1", "Office", 500, "Down", "", 0.0, "", "", "", 0.0, 0.0, "", "", "", ""],
    ]
    raw = build_synthetic_workbook(unit_rows=unit_rows)
    result = normalize_ledger_xlsx(raw)
    assert result["status"] == "normalized_unvalidated"
    assert len(result["residential_units"]) == 1
    assert [u["unit_id"] for u in result["commercial_units"]] == ["R-1", "R-2", "O-1"]
    assert result["counts"]["commercial"] == {
        "occupied": 1,
        "vacant": 1,
        "down": 1,
        "total": 3,
    }


def test_floorplan_does_not_imply_use():
    unit_rows = [
        ["A-1", "1BR", 700, "Occupied", CANARIES[0], 1100.0, "Resident", "RENT", 1100.0, 0.0, 300.0, "01/01/2025", "01/01/2025", "12/31/2025", ""],
        ["A-2", "Studio", 500, "Vacant", "", 900.0, "", "", "", 0.0, 0.0, "", "", "", ""],
        ["A-3", "A1", 650, "Occupied", CANARIES[1], 1000.0, "Resident", "RENT", 1000.0, 0.0, 300.0, "01/01/2025", "01/01/2025", "12/31/2025", ""],
    ]
    raw = build_synthetic_workbook(unit_rows=unit_rows)
    result = normalize_ledger_xlsx(raw)
    assert result["status"] == "blocked"
    assert result["residential_units"] == []
    assert result["commercial_units"] == []
    assert len(unresolved_units(result)) == 3
    assert all(v is None for v in result["counts"]["residential"].values())


def test_missing_status_blocks_without_inventing_counts():
    unit_rows = [
        ["A-1", "Residential", 700, "", CANARIES[0], 1100.0, "Resident", "RENT", 1100.0, 0.0, 300.0, "01/01/2025", "01/01/2025", "12/31/2025", ""],
        ["A-2", "Residential", 700, "Vacant", "", 1100.0, "", "", "", 0.0, 0.0, "", "", "", ""],
    ]
    raw = build_synthetic_workbook(unit_rows=unit_rows)
    result = normalize_ledger_xlsx(raw)
    assert result["status"] == "blocked"
    assert "UNSUPPORTED_STATUS" in codes(result) or "MISSING_UNIT_STATUS" in codes(result)
    assert all(v is None for v in result["counts"]["residential"].values())
    statuses = {u["unit_id"]: u["status"] for u in result["residential_units"]}
    # Missing status must not be coerced into occupied/vacant/down.
    if "A-1" in statuses:
        assert statuses["A-1"] is None


def test_status_mapping():
    unit_rows = [
        ["A-1", "Residential", 700, "Occupied", CANARIES[0], 1100.0, "Resident", "RENT", 1100.0, 0.0, 300.0, "01/01/2025", "01/01/2025", "12/31/2025", ""],
        ["A-2", "Residential", 700, "Notice", CANARIES[1], 1100.0, "Resident", "RENT", 1100.0, 0.0, 300.0, "01/01/2024", "01/01/2024", "12/31/2024", "12/31/2024"],
        ["A-3", "Residential", 700, "Current", CANARIES[2], 1100.0, "Resident", "RENT", 1100.0, 0.0, 300.0, "01/01/2025", "01/01/2025", "12/31/2025", ""],
        ["A-4", "Residential", 700, "Vacant", "", 1100.0, "", "", "", 0.0, 0.0, "", "", "", ""],
        ["A-5", "Residential", 700, "Down", "", 0.0, "", "", "", 0.0, 0.0, "", "", "", ""],
        ["A-6", "Residential", 700, "Offline", "", 0.0, "", "", "", 0.0, 0.0, "", "", "", ""],
    ]
    raw = build_synthetic_workbook(unit_rows=unit_rows)
    result = normalize_ledger_xlsx(raw)
    assert result["counts"]["residential"] == {
        "occupied": 3,
        "vacant": 1,
        "down": 2,
        "total": 6,
    }


def test_admin_down_is_not_independent_down():
    unit_rows = [
        ["A-1", "Residential", 700, "Admin/Down", "", 0.0, "", "", "", 0.0, 0.0, "", "", "", ""],
        ["A-2", "Residential", 700, "Model", "", 0.0, "", "", "", 0.0, 0.0, "", "", "", ""],
    ]
    raw = build_synthetic_workbook(unit_rows=unit_rows)
    result = normalize_ledger_xlsx(raw)
    assert result["status"] == "blocked"
    assert "UNSUPPORTED_STATUS" in codes(result)
    assert all(v is None for v in result["counts"]["residential"].values())
    for unit in result["residential_units"]:
        assert unit["status"] is None


def test_charge_classification_ancillary_before_base_rent():
    assert "lease" not in {c.lower() for c in BASE_RENT_CODES}
    assert "pet rent" in ANCILLARY_CATEGORIES["pet"] or any(
        "pet rent" in codes for codes in ANCILLARY_CATEGORIES.values()
    )
    assert ANCILLARY_CODES

    for code in ("RENT", "Rent", "BASE", "Base Rent", "MRENT", "Market Rent", "Apt Rent", "APTR", "LEASE RENT"):
        assert is_base_rent(code) is True
        assert is_ancillary_charge(code) is False
        assert classify_charge(code) == "base_rent"

    assert classify_charge("PARK") == "parking"
    assert classify_charge("Garage") == "parking"
    assert classify_charge("PET") == "pet"
    assert classify_charge("STOR") == "storage"
    assert classify_charge("UTIL") == "utilities"
    assert classify_charge("W/D") == "washer_dryer"
    assert classify_charge("CONC") == "concessions"


def test_pet_rent_is_pet_not_base_rent():
    assert classify_charge("Pet Rent") == "pet"
    assert is_base_rent("Pet Rent") is False
    assert is_ancillary_charge("Pet Rent") is True


def test_free_rent_is_concessions():
    assert classify_charge("Free Rent") == "concessions"
    assert is_base_rent("Free Rent") is False
    assert is_ancillary_charge("Free Rent") is True


def test_lease_id_start_end_are_not_rent():
    for label in ("Lease ID", "Lease Start", "Lease End", "Lease"):
        assert is_base_rent(label) is False
        assert classify_charge(label) == "unknown"


def test_lease_columns_are_not_cited_as_rent():
    raw = build_synthetic_workbook(
        unit_rows=[["A-1", "Residential", 700, "Occupied", CANARIES[0], 1100.0, "Resident", "RENT", 1100.0, 0.0, 300.0, "01/01/2025", "01/01/2025", "12/31/2025", ""]],
    )
    result = normalize_ledger_xlsx(raw)
    evidence = result["residential_units"][0]["evidence"][0]
    charge_col = evidence["charge"]["column"]
    # Charge Code is column 8; Lease Start/End are 13/14.
    assert charge_col == 8
    assert charge_col not in (12, 13, 14, 15)


def test_citation_coordinate_accuracy():
    raw = build_synthetic_workbook(
        header_row=7,
        unit_rows=[
            ["A-9", "1BR", 650, "Occupied", CANARIES[0], 1100.0, "Resident", "RENT", 1100.0, 0.0, 350.0, "02/01/2025", "02/01/2025", "01/31/2026", ""],
            ["", "", "", "", "", "", "Resident", "PET", 30.0, "", "", "", "", "", ""],
        ],
    )
    result = normalize_ledger_xlsx(raw)
    u = unresolved_units(result)[0]
    assert u["evidence"][0]["unit_id"]["row"] == 8
    assert u["evidence"][0]["unit_id"]["column"] == 1
    assert u["evidence"][0]["unit_id"]["sheet"] == 2
    assert u["evidence"][0]["status"]["column"] == 4
    assert u["evidence"][0]["charge"]["column"] == 8
    assert u["evidence"][1]["charge"]["row"] == 9
    assert u["evidence"][1]["charge"]["column"] == 8
    check_citations(result, raw)


def test_orphan_continuation_emits_blocker():
    unit_rows = [
        ["", "", "", "", "", "", "Resident", "PET", 35.0, "", "", "", "", "", ""],
        ["A-101", "1BR", 700, "Occupied", CANARIES[0], 1200.0, "Resident", "RENT", 1200.0, 0.0, 400.0, "01/01/2025", "01/01/2025", "12/31/2025", ""],
    ]
    raw = build_synthetic_workbook(unit_rows=unit_rows)
    result = normalize_ledger_xlsx(raw)
    assert "ORPHAN_CONTINUATION" in codes(result)
    assert result["status"] == "blocked"


def test_duplicate_unit_conflict_blocks():
    unit_rows = [
        ["A-101", "Residential", 700, "Occupied", CANARIES[0], 1200.0, "Resident", "RENT", 1200.0, 0.0, 400.0, "01/01/2025", "01/01/2025", "12/31/2025", ""],
        ["A-101", "Residential", 700, "Vacant", "", 1200.0, "", "", "", 0.0, 0.0, "", "", "", ""],
    ]
    raw = build_synthetic_workbook(unit_rows=unit_rows)
    result = normalize_ledger_xlsx(raw)
    assert "DUPLICATE_UNIT_CONFLICT" in codes(result)
    assert result["status"] == "blocked"
    assert all(v is None for v in result["counts"]["residential"].values())


def test_future_section_excluded():
    current_units = [
        ["A-101", "1BR", 700, "Occupied", CANARIES[0], 1200.0, "Resident", "RENT", 1200.0, 0.0, 400.0, "01/01/2025", "01/01/2025", "12/31/2025", ""],
    ]
    extra_sections = [
        ("Future Residents", [
            ["A-102", "1BR", 750, "Applicant", "Future Tenant A", 1300.0, "Resident", "RENT", 1300.0, 0.0, 500.0, "03/01/2026", "03/01/2026", "02/28/2027", ""],
        ]),
    ]
    raw = build_synthetic_workbook(unit_rows=current_units, extra_sections=extra_sections)
    result = normalize_ledger_xlsx(raw)
    unresolved = unresolved_units(result)
    assert [u["unit_id"] for u in unresolved] == ["A-101"]
    assert "NON_CURRENT_SECTION_EXCLUDED" in codes(result)
    assert "Future Tenant A" not in json.dumps(result)


def test_totals_terminate_parsing():
    current_units = [
        ["A-101", "1BR", 700, "Occupied", CANARIES[0], 1200.0, "Resident", "RENT", 1200.0, 0.0, 400.0, "01/01/2025", "01/01/2025", "12/31/2025", ""],
    ]
    totals_row = ["Totals", "", "", "", "", 1200.0, "", "", 1200.0]
    raw = build_synthetic_workbook(unit_rows=current_units, totals_row=totals_row)
    unresolved = unresolved_units(normalize_ledger_xlsx(raw))
    assert len(unresolved) == 1
    assert unresolved[0]["unit_id"] == "A-101"


def test_missing_inventory_headers_fails_layout():
    raw = build_synthetic_workbook(
        headers=["Col A", "Col B", "Col C", "Col D"],
        unit_rows=sample_units(),
    )
    with pytest.raises(RentRollNormalizationError) as exc_info:
        normalize_ledger_xlsx(raw)
    assert exc_info.value.code == "UNSUPPORTED_XLSX_LAYOUT"
    assert exc_info.value.__context__ is None
    assert exc_info.value.__cause__ is None


def test_parameters_only_workbook_fails_layout():
    raw = build_synthetic_workbook(
        include_parameters=True,
        headers=["Parameter", "Value"],
        unit_rows=[["As of", "01/15/2026"]],
        inventory_sheet_title="Notes",
    )
    with pytest.raises(RentRollNormalizationError) as exc_info:
        normalize_ledger_xlsx(raw)
    assert exc_info.value.code == "UNSUPPORTED_XLSX_LAYOUT"


@pytest.mark.parametrize("payload", [b"", b"Not a zip file", b"PK\x03\x04corrupt_zip_content", "string_not_bytes", None])
def test_requires_valid_ooxml_bytes(payload):
    with pytest.raises(RentRollNormalizationError) as exc_info:
        normalize_ledger_xlsx(payload)
    assert exc_info.value.code in ("UNSUPPORTED_XLSX_FORMAT", "MALFORMED_INPUT", "EMPTY_INPUT")
    assert exc_info.value.__context__ is None
    assert exc_info.value.__cause__ is None


def test_vba_macro_rejected():
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("xl/vbaProject.bin", b"fake macro content")
        zf.writestr("[Content_Types].xml", b"<xml></xml>")
    with pytest.raises(RentRollNormalizationError) as exc_info:
        normalize_ledger_xlsx(buf.getvalue())
    assert exc_info.value.code == "MALFORMED_INPUT"


def test_external_links_rejected():
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("xl/externalLinks/externalLink1.xml", b"<xml></xml>")
        zf.writestr("[Content_Types].xml", b"<xml></xml>")
    with pytest.raises(RentRollNormalizationError) as exc_info:
        normalize_ledger_xlsx(buf.getvalue())
    assert exc_info.value.code == "MALFORMED_INPUT"


def test_input_byte_limit_exceeded():
    huge_bytes = b"PK\x03\x04" + b"\x00" * (33 * 1024 * 1024)
    with pytest.raises(RentRollNormalizationError) as exc_info:
        normalize_ledger_xlsx(huge_bytes)
    assert exc_info.value.code == "INPUT_LIMIT_EXCEEDED"


def test_column_limit_exceeded():
    raw = build_synthetic_workbook(
        extra_columns=120,
        unit_rows=[["A-1", "Residential", 700, "Occupied", CANARIES[0], 1100.0, "Resident", "RENT", 1100.0]],
    )
    with pytest.raises(RentRollNormalizationError) as exc_info:
        normalize_ledger_xlsx(raw)
    assert exc_info.value.code == "INPUT_LIMIT_EXCEEDED"


def test_sparse_xml_row_rejected_before_allocation():
    raw = build_synthetic_workbook(
        unit_rows=[["A-1", "Residential", 700, "Occupied", CANARIES[0], 1100.0, "Resident", "RENT", 1100.0]],
    )
    src = zipfile.ZipFile(io.BytesIO(raw))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as dest:
        for info in src.infolist():
            data = src.read(info.filename)
            if info.filename.startswith("xl/worksheets/") and b"<sheetData>" in data:
                data = data.replace(
                    b"<sheetData>",
                    b'<sheetData><row r="50001"><c r="A50001" t="inlineStr"><is><t>x</t></is></c></row>',
                    1,
                )
            dest.writestr(info, data)
    with pytest.raises(RentRollNormalizationError) as exc_info:
        normalize_ledger_xlsx(out.getvalue())
    assert exc_info.value.code == "INPUT_LIMIT_EXCEEDED"


def test_missing_dependency_is_lazy_and_sanitized(monkeypatch):
    import plat_harness.ingest.ledger_xlsx as mod

    raw = build_synthetic_workbook(
        unit_rows=[["A-1", "Residential", 700, "Occupied", CANARIES[0], 1100.0, "Resident", "RENT", 1100.0]],
    )
    real_import = __import__

    def fake_import(name, *args, **kwargs):
        if name == "openpyxl" or name.startswith("openpyxl"):
            raise ImportError("missing")
        return real_import(name, *args, **kwargs)

    monkeypatch.setattr(mod, "load_workbook", None)
    monkeypatch.setattr("builtins.__import__", fake_import)
    with pytest.raises(RentRollNormalizationError) as exc_info:
        mod.normalize_ledger_xlsx(raw)
    assert exc_info.value.code == "XLSX_DEPENDENCY_MISSING"
    assert exc_info.value.__context__ is None
    assert exc_info.value.__cause__ is None


def test_library_diagnostics_are_not_exposed(monkeypatch, capsys):
    import plat_harness.ingest.ledger_xlsx as mod

    def boom(*_args, **_kwargs):
        import warnings
        warnings.warn(CANARIES[0])
        raise ValueError(CANARIES[0])

    monkeypatch.setattr(mod, "load_workbook", boom)
    raw = build_synthetic_workbook(
        unit_rows=[["A-1", "Residential", 700, "Occupied", CANARIES[0], 1100.0, "Resident", "RENT", 1100.0]],
    )
    with pytest.raises(RentRollNormalizationError) as exc_info:
        mod.normalize_ledger_xlsx(raw)
    assert exc_info.value.code == "MALFORMED_INPUT"
    assert exc_info.value.__context__ is exc_info.value.__cause__ is None
    assert CANARIES[0] not in "".join(traceback.format_exception(exc_info.value))
    assert CANARIES[0] not in "".join(capsys.readouterr())


def test_repeated_header_band_is_skipped():
    unit_rows = [
        ["A-1", "Residential", 700, "Occupied", CANARIES[0], 1100.0, "Resident", "RENT", 1100.0, 0.0, 300.0, "01/01/2025", "01/01/2025", "12/31/2025", ""],
        list(DEFAULT_HEADERS),
        ["A-2", "Residential", 710, "Vacant", "", 1110.0, "", "", "", 0.0, 0.0, "", "", "", ""],
    ]
    raw = build_synthetic_workbook(unit_rows=unit_rows)
    result = normalize_ledger_xlsx(raw)
    assert [u["unit_id"] for u in result["residential_units"]] == ["A-1", "A-2"]
    assert result["counts"]["residential"]["total"] == 2
