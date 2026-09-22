"""Normalize explicitly supported synthetic-tested PMS export shapes.

No raw rows, workbook metadata, free text, money, or tenant identifiers leave
this module. Output is an unvalidated observation, never underwriting approval.
See docs/PMS_NORMALIZER.md for the deliberately narrow input contract.
"""
from __future__ import annotations

import csv
import hashlib
import io
import re
import warnings
import zipfile
from xml.etree.ElementTree import iterparse

__all__ = ["RentRollNormalizationError", "normalize_rent_roll"]

_MAX_BYTES = 8 * 1024 * 1024
_MAX_XLSX_BYTES = 32 * 1024 * 1024
_MAX_ROWS = 50_000
_MAX_COLUMNS = 128
_MAX_CELLS = 250_000
_COUNTS = ("occupied", "vacant", "down", "total")
_SCOPES = ("residential", "commercial")
_ALIASES = {
    "yardi": "yardi", "yardi voyager": "yardi", "voyager": "yardi",
    "realpage": "realpage", "realpage onesite": "realpage", "onesite": "realpage",
    "entrata": "entrata",
}
# Explicit labels only: a floorplan (e.g. 1BR) does not establish asset use.
_TYPES = {
    "residential": "residential", "apartment": "residential",
    "commercial": "commercial", "retail": "commercial", "office": "commercial",
}
_STATUSES = {
    "occupied": "occupied", "current": "occupied", "notice": "occupied",
    "notice rented": "occupied", "notice unrented": "occupied",
    "vacant": "vacant", "vacant rented": "vacant", "vacant unrented": "vacant",
    "down": "down", "offline": "down", "out of service": "down",
}
_FUTURE_STATUSES = {"applicant", "future", "future resident", "future tenant"}
_FUTURE_SECTIONS = {"future residents", "future tenants", "applicants", "future/applicants"}
_CURRENT_SECTIONS = {"current residents", "current units", "current tenants", "units"}
_TOTALS = {"total", "totals", "grand total", "report total"}
_CONTINUATIONS = {"charge", "co-tenant", "cotenant"}
_UNIT_ID = re.compile(r"(?:[A-Z]{1,3}-?)?[0-9]{1,4}[A-Z]?", re.ASCII)
_CONTEXT_HEADERS = {
    "property": {"property", "property id", "property code", "property name"},
    "building": {"building", "building id", "building number", "building #", "building name"},
}
_COMMON_HEADERS = {
    "tenant_name": {"resident", "resident name", "tenant", "tenant name"},
    "phone": {"phone", "phone number", "telephone"},
    "email": {"email", "email address", "e-mail"},
    "record_type": {"record type", "row type"},
}
_VENDOR_HEADERS = {
    "yardi": {"unit_id": {"unit", "unit number"}, "status": {"status", "unit status"},
              "unit_type": {"unit type", "space type", "property type"}},
    "realpage": {"unit_id": {"unit", "unit #", "unit number"}, "status": {"status", "occupancy status"},
                 "unit_type": {"unit type", "unit category", "space type"}},
    "entrata": {"unit_id": {"unit", "unit number"}, "status": {"status", "occupancy"},
                "unit_type": {"unit type", "space type"}},
}


class RentRollNormalizationError(ValueError):
    """Safe structural-input failure. ``code`` is a static machine-readable tag."""

    def __init__(self, code: str):
        self.code = code
        super().__init__(f"Rent roll normalization failed ({code}).")


class _InputError(Exception):
    """Internal control flow; only static codes are permitted."""


class _FormattedNumeric(str):
    """Internal marker: a numeric cell has non-General display semantics."""


def _xlsx_bounds(sheet):
    """Bound actual coordinates BEFORE openpyxl can allocate sparse row arrays.

    Match inferred row/column counters, but reject malformed/nonmonotone cells
    rather than relying on openpyxl's permissive reordering or overwriting.
    The read-only source hook is isolated here (tested with openpyxl 3.1.5).
    """
    ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    row_number = column = cells = 0
    stack = []
    with sheet._get_source() as source:
        for event, element in iterparse(source, events=("start", "end")):
            if event == "end":
                element.clear()
                stack.pop()
                continue
            parent = stack[-1] if stack else None
            stack.append(element.tag)
            if element.tag == ns + "row":
                raw_row = element.get("r")
                if raw_row is not None and not re.fullmatch(r"[1-9][0-9]*", raw_row, re.ASCII):
                    raise _InputError("MALFORMED_INPUT")
                if raw_row is not None and len(raw_row) > 6:
                    raise _InputError("INPUT_LIMIT_EXCEEDED")
                index = int(raw_row) if raw_row is not None else row_number + 1
                if index > _MAX_ROWS:
                    raise _InputError("INPUT_LIMIT_EXCEEDED")
                if index <= row_number:
                    raise _InputError("MALFORMED_INPUT")
                row_number, column = index, 0
            elif parent == ns + "row":
                if element.tag != ns + "c":
                    raise _InputError("MALFORMED_INPUT")
                coordinate = element.get("r")
                index = column + 1
                if coordinate is not None:
                    match = re.fullmatch(r"([A-Z]+)([1-9][0-9]*)", coordinate, re.ASCII)
                    if match is None:
                        raise _InputError("MALFORMED_INPUT")
                    letters, digits = match.groups()
                    if len(letters) > 3 or len(digits) > 6:
                        raise _InputError("INPUT_LIMIT_EXCEEDED")
                    index = 0
                    for letter in letters:
                        index = index * 26 + ord(letter) - ord("A") + 1
                    if int(digits) > _MAX_ROWS:
                        raise _InputError("INPUT_LIMIT_EXCEEDED")
                    if int(digits) != row_number:
                        raise _InputError("MALFORMED_INPUT")
                if index > _MAX_COLUMNS:
                    raise _InputError("INPUT_LIMIT_EXCEEDED")
                if index <= column:
                    raise _InputError("MALFORMED_INPUT")
                # Padded cells, not merely populated <c> nodes, consume budget.
                cells += index - column
                column = index
                if cells > _MAX_CELLS:
                    raise _InputError("INPUT_LIMIT_EXCEEDED")
    return row_number, cells


def _key(value: str) -> str:
    return " ".join(value.strip().lower().split())


def _cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, (str, int, float)) and not isinstance(value, bool):
        return str(value).strip()
    return "\x00"  # Unsupported cell type, not a representation of the value.


def _tables(raw: bytes):
    """Yield (sheet position, [(physical start, end, cell strings), ...])."""
    if raw.startswith(b"PK"):
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            entries = archive.infolist()
            if (len(entries) > 1024 or sum(i.file_size for i in entries) > _MAX_XLSX_BYTES):
                raise _InputError("INPUT_LIMIT_EXCEEDED")
            if any(i.flag_bits & 1 or "vbaproject" in i.filename.lower()
                           or "externallinks/" in i.filename.lower()
                           or "activex/" in i.filename.lower()
                           or "ctrlprops/" in i.filename.lower()
                           or "embeddings/" in i.filename.lower() for i in entries):
                raise _InputError("MALFORMED_INPUT")
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise _InputError("XLSX_DEPENDENCY_MISSING") from None
        book = load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
        try:
            total_cells = total_rows = 0
            for sheet in book.worksheets:
                if (sheet.max_row or 0) > _MAX_ROWS or (sheet.max_column or 0) > _MAX_COLUMNS:
                    raise _InputError("INPUT_LIMIT_EXCEEDED")
                rows, cells = _xlsx_bounds(sheet)
                total_rows += rows
                total_cells += cells
                if total_rows > _MAX_ROWS or total_cells > _MAX_CELLS:
                    raise _InputError("INPUT_LIMIT_EXCEEDED")
            total_cells = total_rows = 0
            for sheet_number, sheet in enumerate(book.worksheets, start=1):
                # Do not trust producer dimensions to hide rows outside their range.
                sheet.reset_dimensions()
                records = []
                for row_number, row in enumerate(sheet.iter_rows(), start=1):
                    total_rows += 1
                    total_cells += len(row)
                    if total_rows > _MAX_ROWS or total_cells > _MAX_CELLS or len(row) > _MAX_COLUMNS:
                        raise _InputError("INPUT_LIMIT_EXCEEDED")
                    values = [_FormattedNumeric(_cell_text(c.value))
                              if c.data_type == "n" and c.value is not None and c.number_format != "General"
                              else _cell_text(c.value) for c in row]
                    records.append((row_number, row_number, values))
                yield sheet_number, records
        finally:
            book.close()
    else:
        text = raw.decode("utf-8-sig")
        if "\x00" in text:
            raise _InputError("MALFORMED_INPUT")
        reader = csv.reader(io.StringIO(text, newline=""), strict=True)
        records = []
        cells = 0
        start = 1
        for row in reader:
            cells += len(row)
            if reader.line_num > _MAX_ROWS or len(row) > _MAX_COLUMNS or cells > _MAX_CELLS:
                raise _InputError("INPUT_LIMIT_EXCEEDED")
            records.append((start, reader.line_num, [_cell_text(c) for c in row]))
            start = reader.line_num + 1
        yield 1, records


def _header(row, definitions, required):
    found = {}
    duplicates = False
    for index, cell in enumerate(row):
        for field, aliases in definitions.items():
            if _key(cell) in aliases:
                if field in found:
                    duplicates = True
                found[field] = index
    if not set(required).issubset(found):
        return None
    if duplicates:
        raise _InputError("AMBIGUOUS_HEADER")
    return found


def _normalize(raw: bytes, pms: str) -> dict:
    digest = hashlib.sha256(raw).hexdigest()
    result = {
        "version": "1.0", "pms_type": pms, "source_sha256": digest,
        "residential_units": [], "commercial_units": [],
        "counts": {scope: dict.fromkeys(_COUNTS) for scope in _SCOPES},
        "summary": {scope: {"status": "absent", "reported_counts": None, "citations": []}
                    for scope in _SCOPES},
        "issues": [], "status": "normalized_unvalidated",
    }
    inventory_valid = True
    has_header = False
    observed = {}
    conflicts = set()
    context = None
    context_valid = True
    definitions = {**_COMMON_HEADERS, **_VENDOR_HEADERS[pms], **_CONTEXT_HEADERS}
    summary_definitions = {field: {field} for field in _COUNTS}
    summary_definitions["scope"] = {"scope", "summary scope"}

    def citation(sheet, start, end, column):
        return {"source_sha256": digest, "sheet": sheet, "row": start,
                "row_end": end, "column": column + 1}

    def issue(code, cite=None, *, inventory=False, warning=False):
        nonlocal inventory_valid
        result["issues"].append({"code": code, "severity": "warning" if warning else "blocker",
                                 "citation": cite})
        if not warning:
            result["status"] = "blocked"
        if inventory:
            inventory_valid = False

    for sheet, records in _tables(raw):
        mode = None
        header = None
        width = 0
        skip_future = False
        sheet_header = False
        for start, end, row in records:
            if not any(row):
                continue
            cite = citation(sheet, start, end, 0)
            first = _key(row[0]) if row else ""
            singleton = not any(row[1:])
            if singleton and first in _FUTURE_SECTIONS:
                skip_future = True
                issue("FUTURE_OR_APPLICANT_EXCLUDED", cite, warning=True)
                continue
            if singleton and first in _CURRENT_SECTIONS:
                skip_future = False
                continue
            new_summary_header = _header(row, summary_definitions, (*_COUNTS, "scope"))
            if new_summary_header:
                mode, header, width = "summary", new_summary_header, len(row)
                continue
            new_header = _header(row, definitions, ("unit_id", "status", "unit_type"))
            if new_header:
                has_header = sheet_header = True
                mode, header, width = "units", new_header, len(row)
                continue
            fields = {field for field, aliases in definitions.items()
                      if any(_key(cell) in aliases for cell in row)}
            if (("unit_id" in fields and (mode is None or len(fields) > 1 or singleton))
                    or {"status", "unit_type"}.issubset(fields)):
                sheet_header = True  # Not an arbitrary metadata sheet.
                issue("INCOMPLETE_INVENTORY_HEADER", cite, inventory=True)
                mode, header = None, None
                continue
            if mode is None:
                continue  # Preamble is never returned.
            if mode == "summary":
                values = {k: row[i] if i < len(row) else "" for k, i in header.items()}
                scope = _key(values["scope"])
                scope = {"residential summary": "residential", "commercial summary": "commercial"}.get(scope, scope)
                if scope not in _SCOPES:
                    issue("UNSUPPORTED_SUMMARY_SCOPE", cite)
                    continue
                summary = result["summary"][scope]
                cites = {k: citation(sheet, start, end, header[k]) for k in _COUNTS}
                summary["citations"].append(cites)
                if (any(not re.fullmatch(r"[0-9]{1,6}", values[k], re.ASCII) for k in _COUNTS)
                        or any(row[width:])):
                    summary["status"], summary["reported_counts"] = "invalid", None
                    issue("INVALID_SUMMARY", cite)
                    continue
                reported = {k: int(values[k]) for k in _COUNTS}
                if summary["status"] == "invalid":
                    continue  # Never heal an invalid observation with a later row.
                if summary["reported_counts"] is not None:
                    if summary["reported_counts"] != reported:
                        summary["status"], summary["reported_counts"] = "invalid", None
                        issue("SUMMARY_CONFLICT", cite)
                    else:
                        issue("DUPLICATE_SUMMARY_EVIDENCE", cite, warning=True)
                    continue
                summary["reported_counts"] = reported
                summary["status"] = "unresolved"
                continue
            if skip_future:
                continue
            if first in _TOTALS:
                issue("UNSCOPED_REPORT_TOTAL", cite)
                continue
            if max(header[k] for k in ("unit_id", "status", "unit_type")) >= len(row) or any(row[width:]):
                issue("MALFORMED_ROW", cite, inventory=True)
                continue
            values = {k: row[i] if i < len(row) else "" for k, i in header.items()}
            status_text = _key(values["status"])
            role = _key(values.get("record_type", "")) or "unit"
            if status_text in _FUTURE_STATUSES or role in _FUTURE_STATUSES:
                issue("FUTURE_OR_APPLICANT_EXCLUDED", cite, warning=True)
                continue
            if role not in {"unit", *_CONTINUATIONS}:
                issue("UNSUPPORTED_RECORD_TYPE", cite, inventory=True)
                continue
            row_context = tuple(values.get(field) for field in _CONTEXT_HEADERS)
            if any(field in header and (not values[field] or values[field].startswith("=")
                                       or "\x00" in values[field]
                                       or isinstance(values[field], _FormattedNumeric))
                   for field in _CONTEXT_HEADERS):
                context_valid = False
                issue("INVALID_UNIT_CONTEXT", cite, inventory=True)
                continue
            if context is None:
                context = row_context
            elif context != row_context:
                context_valid = False
                issue("AMBIGUOUS_UNIT_CONTEXT", cite, inventory=True)
                continue
            unit_id = values["unit_id"]
            if isinstance(unit_id, _FormattedNumeric):
                issue("UNSUPPORTED_UNIT_ID_FORMAT", citation(sheet, start, end, header["unit_id"]),
                      inventory=True)
                continue
            if not _UNIT_ID.fullmatch(unit_id):
                issue("INVALID_UNIT_ID", citation(sheet, start, end, header["unit_id"]), inventory=True)
                continue
            if unit_id in conflicts:
                issue("DUPLICATE_UNIT_CONFLICT", cite, inventory=True)
                continue
            existing = observed.get(unit_id)
            status = _STATUSES.get(status_text)
            unit_type = _TYPES.get(_key(values["unit_type"]))
            if role in _CONTINUATIONS:
                if existing is None:
                    issue("ORPHAN_CONTINUATION", cite, inventory=True)
                    continue
                if not values["status"]:
                    status = existing["status"]
                if not values["unit_type"]:
                    unit_type = existing["unit_type"]
            if status is None:
                issue("UNSUPPORTED_STATUS", citation(sheet, start, end, header["status"]), inventory=True)
            if unit_type is None:
                issue("UNSUPPORTED_UNIT_TYPE", citation(sheet, start, end, header["unit_type"]), inventory=True)
            if status is None or unit_type is None:
                continue
            if existing and (existing["status"] != status or existing["unit_type"] != unit_type):
                conflicts.add(unit_id)
                del observed[unit_id]
                issue("DUPLICATE_UNIT_CONFLICT", cite, inventory=True)
                continue
            evidence = {k: citation(sheet, start, end, i) for k, i in header.items()
                        if k not in _CONTEXT_HEADERS and i < len(row) and values[k]}
            redacted = {k: "[REDACTED]" for k in ("tenant_name", "phone", "email") if k in header}
            if existing:
                existing["evidence"].append(evidence)
                existing.update(redacted)
                issue("CONTINUATION_EVIDENCE" if role in _CONTINUATIONS else "DUPLICATE_UNIT_EVIDENCE",
                      cite, warning=True)
            else:
                observed[unit_id] = {"unit_id": unit_id, "status": status, "unit_type": unit_type,
                                     **redacted, "evidence": [evidence]}
        if records and not sheet_header:
            issue("NON_INVENTORY_SHEET_IGNORED", warning=True)
    if not has_header:
        raise _InputError("HEADER_NOT_FOUND")
    if not context_valid:
        observed.clear()  # Do not expose an apparently deduplicated partial inventory.
    if not observed:
        issue("NO_CURRENT_UNITS", inventory=True)
    for unit in observed.values():
        result[unit["unit_type"] + "_units"].append(unit)
    if inventory_valid:
        for scope in _SCOPES:
            units = result[scope + "_units"]
            counts = {state: sum(unit["status"] == state for unit in units) for state in _COUNTS[:-1]}
            counts["total"] = len(units)
            result["counts"][scope] = counts
    for scope in _SCOPES:
        summary = result["summary"][scope]
        if summary["status"] == "unresolved" and inventory_valid:
            if summary["reported_counts"] == result["counts"][scope]:
                summary["status"] = "reconciled"
            else:
                summary["status"] = "mismatch"
                issue("SUMMARY_MISMATCH", summary["citations"][0]["total"])
    return result


def normalize_rent_roll(stream, pms_type) -> dict:
    """Consume CSV/XLSX or supported RealPage BIFF from its current position.

    Return JSON-safe, allowlisted observations and content-bound cell citations.
    Structural failures raise only a sanitized RentRollNormalizationError;
    semantic uncertainty returns status='blocked' with static issue codes.
    """
    error_code = None
    try:
        # Library warnings can interpolate workbook metadata. Turn them into
        # sanitized structural failures before a warnings/logging sink sees them.
        with warnings.catch_warnings():
            warnings.simplefilter("error")
            pms = _ALIASES.get(_key(pms_type)) if isinstance(pms_type, str) else None
            if pms is None:
                raise _InputError("UNSUPPORTED_PMS")
            payload = stream.read(_MAX_BYTES + 1)
            if isinstance(payload, str):
                raw = payload.encode("utf-8")
            elif isinstance(payload, bytes):
                raw = payload
            else:
                raise _InputError("MALFORMED_INPUT")
            if len(raw) > _MAX_BYTES:
                raise _InputError("INPUT_LIMIT_EXCEEDED")
            if not raw.strip():
                raise _InputError("EMPTY_INPUT")
            if raw.startswith(b"%PDF-"):
                raise _InputError("UNSUPPORTED_INPUT_FORMAT")
            if raw.startswith(bytes.fromhex("d0cf11e0a1b11ae1")):
                if pms != "realpage":
                    raise _InputError("UNSUPPORTED_PMS_FORMAT")
                from .onesite import normalize_onesite_xls
                try:
                    return normalize_onesite_xls(raw)
                except RentRollNormalizationError as exc:
                    raise _InputError(exc.code) from None
            return _normalize(raw, pms)
    except _InputError as exc:
        error_code = exc.args[0]
    except Exception:
        error_code = "MALFORMED_INPUT"
    # Raise OUTSIDE handlers: even __context__ cannot retain an input-bearing
    # decoder/library/stream exception for a downstream logger to serialize.
    raise RentRollNormalizationError(error_code) from None
