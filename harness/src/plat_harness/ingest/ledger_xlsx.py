"""Read-only ledger/charge-detail XLSX adapter; never occupancy approval.

Unknown-vendor OOXML workbooks with a physical-unit header band
(Bldg-Unit, Unit Status, Ledger, Charge Code, Scheduled Charges) and a
separate non-inventory parameters sheet. Inventory sheets are selected by
headers, never by worksheet title.
No raw rows, workbook metadata, free text, money, or tenant identifiers leave
this module. Output is an unvalidated observation, never underwriting approval.
"""
from __future__ import annotations

import hashlib
import io
import re
import warnings
import zipfile
from typing import Any
from xml.etree.ElementTree import iterparse

from .pms_normalizer import RentRollNormalizationError

__all__ = [
    "ANCILLARY_CATEGORIES",
    "ANCILLARY_CODES",
    "BASE_RENT_CODES",
    "classify_charge",
    "is_ancillary_charge",
    "is_base_rent",
    "normalize_ledger_xlsx",
]

load_workbook = None

_MAX_BYTES = 32 * 1024 * 1024
_MAX_ROWS = 50_000
_MAX_COLUMNS = 128
_MAX_CELLS = 250_000

_COUNTS = ("occupied", "vacant", "down", "total")
_SCOPES = ("residential", "commercial")

_TYPES = {
    "residential": "residential",
    "apartment": "residential",
    "commercial": "commercial",
    "retail": "commercial",
    "office": "commercial",
}

_STATUSES = {
    "occupied": "occupied",
    "current": "occupied",
    "notice": "occupied",
    "notice rented": "occupied",
    "notice unrented": "occupied",
    "vacant": "vacant",
    "vacant rented": "vacant",
    "vacant unrented": "vacant",
    "vacant leased": "vacant",
    "vacant not leased": "vacant",
    "down": "down",
    "offline": "down",
    "out of service": "down",
}

BASE_RENT_CODES = frozenset({
    "rent",
    "base",
    "base rent",
    "baserent",
    "mrent",
    "market rent",
    "apt rent",
    "aptr",
    "lease rent",
    "resrent",
    "contract rent",
})

_NON_CHARGE_HEADERS = frozenset({
    "lease",
    "lease id",
    "lease start",
    "lease end",
    "lease term",
    "resh id",
    "resident id",
    "tenant id",
    "move in",
    "move out",
    "expected move out",
    "balance",
    "deposit",
    "deposit held",
    "bldg unit",
    "unit type",
    "unit status",
    "resident",
    "ledger",
})

ANCILLARY_CATEGORIES: dict[str, frozenset[str]] = {
    "parking": frozenset({
        "park", "parking", "garage", "carport", "space", "lot",
        "covered parking", "reserved parking",
    }),
    "pet": frozenset({
        "pet", "pet rent", "petrent", "pet fee", "animal", "dog", "cat",
    }),
    "storage": frozenset({
        "stor", "storage", "locker", "storage locker",
    }),
    "utilities": frozenset({
        "util", "utility", "water", "sewer", "trash", "electric", "gas",
        "rubs", "cable", "tech", "valet trash", "pest", "cam",
    }),
    "washer_dryer": frozenset({
        "washer", "dryer", "w/d", "wd", "w d", "laundry", "appliance",
    }),
    "concessions": frozenset({
        "conc", "concession", "discount", "free rent", "credit",
    }),
}

ANCILLARY_CODES = frozenset.union(*ANCILLARY_CATEGORIES.values())

_EXCLUDED_SECTIONS = frozenset({
    "future residents",
    "future tenants",
    "applicants",
    "future/applicants",
    "future residents/applicants",
    "pending renewals",
    "former residents",
    "future",
    "applicant",
})

_TOTALS = frozenset({
    "total",
    "totals",
    "totals:",
    "grand total",
    "report total",
    "summary",
})

_HEADER_ALIASES = {
    "bldg_unit": frozenset({"bldg unit", "building unit", "bldg unit id", "building unit id"}),
    "unit_type": frozenset({"unit type", "type", "space type", "property type", "use"}),
    "sqft": frozenset({"sqft", "sq ft", "square feet", "sq. ft."}),
    "status": frozenset({"unit status", "status", "occupancy status"}),
    "resident": frozenset({"resident", "resident name", "tenant", "tenant name", "name"}),
    "market_rent": frozenset({"market rent", "market"}),
    "ledger": frozenset({"ledger"}),
    "charge_code": frozenset({"charge code", "charge"}),
    "scheduled": frozenset({"scheduled charges", "scheduled charge", "scheduled"}),
    "balance": frozenset({"balance"}),
    "deposit": frozenset({"deposit held", "deposit"}),
    "move_in": frozenset({"move in"}),
    "lease_start": frozenset({"lease start"}),
    "lease_end": frozenset({"lease end"}),
    "lease_id": frozenset({"lease id"}),
    "move_out": frozenset({"expected move out", "move out"}),
}

_REQUIRED_HEADERS = ("bldg_unit", "status", "charge_code", "scheduled")


class _Failure(Exception):
    """Internal static code failure."""


def _key(value: Any) -> str:
    if value is None:
        return ""
    text = value if isinstance(value, str) else str(value)
    return re.sub(r"[\s/,:\\_.-]+", " ", text.strip().lower()).strip()


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (str, int, float)) and not isinstance(value, bool):
        return str(value).strip()
    return ""


def classify_charge(code: str | None) -> str:
    """Classify charge code; ancillary categories win over base rent."""
    if not code:
        return "unknown"
    normalized = _key(code)
    if not normalized:
        return "unknown"
    if normalized in _TOTALS or normalized in _NON_CHARGE_HEADERS:
        return "unknown"
    for category, codes in ANCILLARY_CATEGORIES.items():
        if normalized in codes or any(
            c == normalized or normalized.startswith(c + " ") or normalized.endswith(" " + c)
            for c in codes
        ):
            return category
    if normalized in BASE_RENT_CODES or any(
        normalized == item or normalized.startswith(item + " ") or normalized.endswith(" " + item)
        for item in BASE_RENT_CODES
    ):
        return "base_rent"
    return "unknown"


def is_base_rent(code: str | None) -> bool:
    return classify_charge(code) == "base_rent"


def is_ancillary_charge(code: str | None) -> bool:
    cat = classify_charge(code)
    return cat not in ("base_rent", "unknown")


def _header_map(row: list[str]) -> dict[str, int] | None:
    found: dict[str, int] = {}
    for index, cell in enumerate(row):
        key = _key(cell)
        if not key:
            continue
        for field, aliases in _HEADER_ALIASES.items():
            if key in aliases and field not in found:
                found[field] = index
                break
    if not all(field in found for field in _REQUIRED_HEADERS):
        return None
    return found


def _xlsx_bounds(sheet: Any) -> tuple[int, int]:
    """Bound actual coordinates BEFORE openpyxl can allocate sparse row arrays."""
    ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    row_number = column = cells = 0
    stack: list[str] = []
    with sheet._get_source() as source:
        for event, element in iterparse(source, events=("start", "end")):
            if event == "end":
                element.clear()
                stack.pop()
                continue
            parent = stack[-1] if stack else None
            stack.append(element.tag)
            if element.tag == ns + "row":
                raw_row = element.get("r")
                if raw_row is not None and not re.fullmatch(r"[1-9][0-9]*", raw_row, re.ASCII):
                    raise _Failure("MALFORMED_INPUT")
                if raw_row is not None and len(raw_row) > 6:
                    raise _Failure("INPUT_LIMIT_EXCEEDED")
                index = int(raw_row) if raw_row is not None else row_number + 1
                if index > _MAX_ROWS:
                    raise _Failure("INPUT_LIMIT_EXCEEDED")
                if index <= row_number:
                    raise _Failure("MALFORMED_INPUT")
                row_number, column = index, 0
            elif parent == ns + "row":
                if element.tag != ns + "c":
                    raise _Failure("MALFORMED_INPUT")
                coordinate = element.get("r")
                index = column + 1
                if coordinate is not None:
                    match = re.fullmatch(r"([A-Z]+)([1-9][0-9]*)", coordinate, re.ASCII)
                    if match is None:
                        raise _Failure("MALFORMED_INPUT")
                    letters, digits = match.groups()
                    if len(letters) > 3 or len(digits) > 6:
                        raise _Failure("INPUT_LIMIT_EXCEEDED")
                    index = 0
                    for letter in letters:
                        index = index * 26 + ord(letter) - ord("A") + 1
                    if int(digits) > _MAX_ROWS:
                        raise _Failure("INPUT_LIMIT_EXCEEDED")
                    if int(digits) != row_number:
                        raise _Failure("MALFORMED_INPUT")
                if index > _MAX_COLUMNS:
                    raise _Failure("INPUT_LIMIT_EXCEEDED")
                if index <= column:
                    raise _Failure("MALFORMED_INPUT")
                cells += index - column
                column = index
                if cells > _MAX_CELLS:
                    raise _Failure("INPUT_LIMIT_EXCEEDED")
    return row_number, cells


def _open_book(raw: bytes) -> Any:
    opener = load_workbook
    if opener is None:
        try:
            from openpyxl import load_workbook as opener
        except ImportError:
            raise _Failure("XLSX_DEPENDENCY_MISSING") from None
    return opener(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)


def _inspect_zip(raw: bytes) -> None:
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        entries = archive.infolist()
        if len(entries) > 1024 or sum(i.file_size for i in entries) > _MAX_BYTES:
            raise _Failure("INPUT_LIMIT_EXCEEDED")
        if any(
            i.flag_bits & 1
            or "vbaproject" in i.filename.lower()
            or "externallinks/" in i.filename.lower()
            or "activex/" in i.filename.lower()
            or "ctrlprops/" in i.filename.lower()
            or "embeddings/" in i.filename.lower()
            for i in entries
        ):
            raise _Failure("MALFORMED_INPUT")


def _parse_workbook(raw: bytes) -> dict[str, Any]:
    _inspect_zip(raw)
    book = _open_book(raw)
    try:
        total_rows = total_cells = 0
        sheets: list[tuple[int, Any, list[list[str]]]] = []
        for idx, ws in enumerate(book.worksheets, start=1):
            if (ws.max_row or 0) > _MAX_ROWS or (ws.max_column or 0) > _MAX_COLUMNS:
                raise _Failure("INPUT_LIMIT_EXCEEDED")
            rows, cells = _xlsx_bounds(ws)
            total_rows += rows
            total_cells += cells
            if total_rows > _MAX_ROWS or total_cells > _MAX_CELLS:
                raise _Failure("INPUT_LIMIT_EXCEEDED")
            ws.reset_dimensions()
            rows_data: list[list[str]] = []
            sheet_cells = 0
            for row in ws.iter_rows(values_only=True):
                if len(row) > _MAX_COLUMNS:
                    raise _Failure("INPUT_LIMIT_EXCEEDED")
                sheet_cells += len(row)
                if sheet_cells > _MAX_CELLS or len(rows_data) + 1 > _MAX_ROWS:
                    raise _Failure("INPUT_LIMIT_EXCEEDED")
                rows_data.append([_cell_text(v) for v in row])
            sheets.append((idx, ws, rows_data))

        target_idx = None
        target_rows: list[list[str]] | None = None
        header_row_idx = None
        col_map = None
        for idx, _ws, rows_data in sheets:
            for r_idx, row in enumerate(rows_data):
                mapped = _header_map(row)
                if mapped is not None:
                    target_idx = idx
                    target_rows = rows_data
                    header_row_idx = r_idx
                    col_map = mapped
                    break
            if target_idx is not None:
                break

        if target_idx is None or target_rows is None or header_row_idx is None or col_map is None:
            raise _Failure("UNSUPPORTED_XLSX_LAYOUT")

        digest = hashlib.sha256(raw).hexdigest()
        result: dict[str, Any] = {
            "version": "1.0",
            "pms_type": "unknown",
            "source_sha256": digest,
            "residential_units": [],
            "commercial_units": [],
            "counts": {scope: dict.fromkeys(_COUNTS) for scope in _SCOPES},
            "summary": {
                scope: {"status": "absent", "reported_counts": None, "citations": []}
                for scope in _SCOPES
            },
            "issues": [],
            "status": "normalized_unvalidated",
        }

        def cite(r: int, c: int) -> dict[str, Any]:
            return {
                "source_sha256": digest,
                "sheet": target_idx,
                "row": r,
                "row_end": r,
                "column": c + 1,
            }

        def issue(
            code: str,
            r: int | None = None,
            c: int = 0,
            *,
            warning: bool = False,
            observation: dict[str, Any] | None = None,
        ) -> None:
            entry: dict[str, Any] = {
                "code": code,
                "severity": "warning" if warning else "blocker",
                "citation": cite(r, c) if r is not None else None,
            }
            if observation is not None:
                entry["observation"] = observation
            result["issues"].append(entry)
            if not warning:
                result["status"] = "blocked"

        if len(sheets) > 1:
            issue("NON_INVENTORY_SHEET_IGNORED", warning=True)

        units: dict[str, dict[str, Any]] = {}
        conflicts: set[str] = set()
        current_unit: dict[str, Any] | None = None
        current_unit_id: str | None = None
        inventory_valid = True

        col_unit = col_map["bldg_unit"]
        col_status = col_map["status"]
        col_charge = col_map["charge_code"]
        col_type = col_map.get("unit_type")
        col_name = col_map.get("resident")
        col_ledger = col_map.get("ledger")
        col_scheduled = col_map.get("scheduled")

        def _at(row: list[str], index: int | None) -> str:
            if index is None or index >= len(row):
                return ""
            return row[index]

        for r_idx in range(header_row_idx + 1, len(target_rows)):
            row = target_rows[r_idx]
            sheet_row = r_idx + 1
            if not any(row):
                continue
            if _header_map(row) is not None:
                continue

            first_text = _key(row[0]) if row else ""
            row_keys = [_key(cell) for cell in row if cell]
            if first_text in _TOTALS or any(k in _TOTALS or k.startswith("total") for k in row_keys[:2]):
                break
            if first_text in _EXCLUDED_SECTIONS or any(
                k in _EXCLUDED_SECTIONS or "future" in k or "applicant" in k for k in row_keys[:2]
            ):
                issue("NON_CURRENT_SECTION_EXCLUDED", sheet_row, 0, warning=True)
                break

            unit_val = _at(row, col_unit)
            type_val = _at(row, col_type)
            name_val = _at(row, col_name)
            status_val = _at(row, col_status)
            charge_val = _at(row, col_charge)
            ledger_val = _at(row, col_ledger)
            scheduled_val = _at(row, col_scheduled)
            has_charge = bool(charge_val or scheduled_val or ledger_val)

            is_continuation = False
            if not unit_val and has_charge:
                is_continuation = True
            elif (
                unit_val
                and current_unit_id is not None
                and unit_val == current_unit_id
                and not status_val
                and not type_val
                and has_charge
            ):
                is_continuation = True

            if is_continuation:
                if current_unit is None:
                    inventory_valid = False
                    issue("ORPHAN_CONTINUATION", sheet_row, col_charge)
                else:
                    evidence_item: dict[str, Any] = {}
                    if charge_val:
                        evidence_item["charge"] = cite(sheet_row, col_charge)
                    elif scheduled_val and col_scheduled is not None:
                        evidence_item["charge"] = cite(sheet_row, col_scheduled)
                    current_unit["evidence"].append(evidence_item or {"charge": cite(sheet_row, col_charge)})
                    issue("CONTINUATION_EVIDENCE", sheet_row, col_charge, warning=True)
                continue

            if not unit_val:
                continue

            unit_id = unit_val
            status_key = _key(status_val)
            if not status_val:
                unit_status = None
                inventory_valid = False
                issue("MISSING_UNIT_STATUS", sheet_row, col_status)
            elif status_key in _STATUSES:
                unit_status = _STATUSES[status_key]
            else:
                unit_status = None
                inventory_valid = False
                issue("UNSUPPORTED_STATUS", sheet_row, col_status)

            type_key = _key(type_val)
            resolved_use = _TYPES.get(type_key)

            evidence_primary: dict[str, Any] = {
                "unit_id": cite(sheet_row, col_unit),
                "status": cite(sheet_row, col_status),
            }
            if type_val and col_type is not None:
                evidence_primary["unit_type"] = cite(sheet_row, col_type)
            if name_val and unit_status == "occupied" and col_name is not None:
                evidence_primary["tenant_name"] = cite(sheet_row, col_name)
            if charge_val:
                evidence_primary["charge"] = cite(sheet_row, col_charge)

            if unit_id in conflicts:
                issue("DUPLICATE_UNIT_CONFLICT", sheet_row, col_unit)
                current_unit = None
                current_unit_id = None
                continue

            existing = units.get(unit_id)
            if existing:
                if existing["status"] != unit_status:
                    inventory_valid = False
                    conflicts.add(unit_id)
                    existing["status"] = None
                    issue("DUPLICATE_UNIT_CONFLICT", sheet_row, col_unit)
                else:
                    issue("CONTINUATION_EVIDENCE", sheet_row, col_charge, warning=True)
                existing["evidence"].append(evidence_primary)
                current_unit = existing
                current_unit_id = unit_id
            else:
                current_unit = {
                    "unit_id": unit_id,
                    "status": unit_status,
                    "unit_type": resolved_use,
                    "tenant_name": "[REDACTED]",
                    "evidence": [evidence_primary],
                }
                units[unit_id] = current_unit
                current_unit_id = unit_id

        if not units:
            inventory_valid = False
            issue("NO_CURRENT_UNITS", header_row_idx + 1, 0)

        has_unresolved = False
        for unit in units.values():
            if unit["unit_type"] is None:
                has_unresolved = True
                u_row = unit["evidence"][0]["unit_id"]["row"]
                issue("UNRESOLVED_UNIT_USE", u_row, col_type or 0, observation=unit)
            elif unit["unit_type"] == "residential":
                result["residential_units"].append(unit)
            elif unit["unit_type"] == "commercial":
                result["commercial_units"].append(unit)

        if has_unresolved or not inventory_valid:
            result["status"] = "blocked"
            for scope in _SCOPES:
                result["counts"][scope] = dict.fromkeys(_COUNTS)
        else:
            for scope in _SCOPES:
                scope_units = result[f"{scope}_units"]
                result["counts"][scope] = {
                    "occupied": sum(u["status"] == "occupied" for u in scope_units),
                    "vacant": sum(u["status"] == "vacant" for u in scope_units),
                    "down": sum(u["status"] == "down" for u in scope_units),
                    "total": len(scope_units),
                }
        return result
    finally:
        book.close()


def normalize_ledger_xlsx(raw_bytes: bytes) -> dict[str, Any]:
    """Parse ledger/charge-detail XLSX into a v1 observation dict.

    Never executes formulas, external links, or macros.
    Tenant identities are strictly redacted. Vendor remains unknown.
    """
    code = "MALFORMED_INPUT"
    try:
        with warnings.catch_warnings():
            warnings.simplefilter("error")
            if not isinstance(raw_bytes, bytes):
                raise _Failure("MALFORMED_INPUT")
            if len(raw_bytes) > _MAX_BYTES:
                raise _Failure("INPUT_LIMIT_EXCEEDED")
            if not raw_bytes:
                raise _Failure("EMPTY_INPUT")
            if not raw_bytes.startswith(b"PK"):
                raise _Failure("UNSUPPORTED_XLSX_FORMAT")
            return _parse_workbook(raw_bytes)
    except _Failure as exc:
        code = exc.args[0] if exc.args else "MALFORMED_INPUT"
    except Exception:
        code = "MALFORMED_INPUT"
    raise RentRollNormalizationError(code) from None
