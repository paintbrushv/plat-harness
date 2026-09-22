"""Read-only split-header / lease-charge XLSX adapter; never occupancy approval.

Normalized extraction for Abilene and similar asset-level rent roll workbooks with
split column headers and charge continuation rows.
No raw rows, workbook metadata, free text, money, or tenant identifiers leave
this module. Output is an unvalidated observation, never underwriting approval.
"""
from __future__ import annotations

import hashlib
import io
import re
import warnings
import zipfile
from typing import Any

from .pms_normalizer import RentRollNormalizationError

__all__ = [
    "BASE_RENT_CODES",
    "ANCILLARY_CATEGORIES",
    "ANCILLARY_CODES",
    "classify_charge",
    "is_base_rent",
    "is_ancillary_charge",
    "normalize_lease_charge_xlsx",
]

_MAX_BYTES = 32 * 1024 * 1024
_MAX_ROWS = 50_000
_MAX_COLUMNS = 128
_MAX_CELLS = 250_000

_COUNTS = ("occupied", "vacant", "down", "total")
_SCOPES = ("residential", "commercial")

# Explicit use labels only: floorplans (e.g. 1BR) do not establish asset use.
_TYPES = {
    "residential": "residential",
    "apartment": "residential",
    "commercial": "commercial",
    "retail": "commercial",
    "office": "commercial",
}

_STATUSES = {
    "occupied": "occupied",
    "current": "occupied",
    "notice": "occupied",
    "notice rented": "occupied",
    "notice unrented": "occupied",
    "vacant": "vacant",
    "vacant rented": "vacant",
    "vacant unrented": "vacant",
    "down": "down",
    "offline": "down",
    "out of service": "down",
}

BASE_RENT_CODES = frozenset({
    "rent",
    "base",
    "base rent",
    "baserent",
    "mrent",
    "market rent",
    "apt rent",
    "aptr",
    "lease",
    "lease rent",
    "resrent",
    "contract rent",
})

ANCILLARY_CATEGORIES: dict[str, frozenset[str]] = {
    "parking": frozenset({"park", "parking", "garage", "carport", "space", "lot", "covered parking", "reserved parking"}),
    "pet": frozenset({"pet", "pet rent", "petrent", "pet fee", "animal", "dog", "cat"}),
    "storage": frozenset({"stor", "storage", "locker", "storage locker"}),
    "utilities": frozenset({"util", "utility", "water", "sewer", "trash", "electric", "gas", "rubs", "cable", "tech", "valet trash", "pest"}),
    "washer_dryer": frozenset({"washer", "dryer", "w/d", "wd", "w d", "laundry", "appliance"}),
    "concessions": frozenset({"conc", "concession", "discount", "free rent", "credit"}),
}

ANCILLARY_CODES = frozenset.union(*ANCILLARY_CATEGORIES.values())

_EXCLUDED_SECTIONS = frozenset({
    "future residents",
    "future tenants",
    "applicants",
    "future/applicants",
    "pending renewals",
    "former residents",
    "future",
    "applicant",
})

_TOTALS = frozenset({
    "total",
    "totals",
    "grand total",
    "report total",
    "summary",
})


class _Failure(Exception):
    """Internal static code failure."""


def _key(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return re.sub(r"[\s/,:\-_]+", " ", value.strip().lower()).strip()
    return re.sub(r"[\s/,:\-_]+", " ", str(value).strip().lower()).strip()


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (str, int, float)) and not isinstance(value, bool):
        return str(value).strip()
    return ""


def classify_charge(code: str | None) -> str:
    """Classify charge code into 'base_rent', specific ancillary category, or 'other_ancillary'."""
    if not code:
        return "unknown"
    normalized = _key(code)
    if normalized in BASE_RENT_CODES:
        return "base_rent"
    for category, codes in ANCILLARY_CATEGORIES.items():
        if normalized in codes or any(c == normalized or normalized.startswith(c + " ") or normalized.endswith(" " + c) for c in codes):
            return category
    return "other_ancillary"


def is_base_rent(code: str | None) -> bool:
    return classify_charge(code) == "base_rent"


def is_ancillary_charge(code: str | None) -> bool:
    cat = classify_charge(code)
    return cat != "base_rent" and cat != "unknown"


def _fold_headers(row5: list[str], row6: list[str]) -> dict[str, int]:
    """Fold split rows 5 and 6 into logical column map."""
    max_len = max(len(row5), len(row6))
    folded_cols: list[str] = []
    for c in range(max_len):
        v5 = row5[c] if c < len(row5) else ""
        v6 = row6[c] if c < len(row6) else ""
        text = f"{v5} {v6}".strip()
        folded_cols.append(_key(text))

    column_map: dict[str, int] = {}

    def match_alias(index: int, aliases: frozenset[str]) -> bool:
        k = folded_cols[index]
        return k in aliases

    # Match each logical column
    for c, text in enumerate(folded_cols):
        if not text:
            continue
        if "unit" not in column_map and text in {"unit", "unit number", "unit #"}:
            column_map["unit"] = c
        elif "unit_type" not in column_map and text in {"unit type", "type", "unit category"}:
            column_map["unit_type"] = c
        elif "sq_ft" not in column_map and ("sq ft" in text or "sqft" in text or "square feet" in text):
            column_map["sq_ft"] = c
        elif "market_rent" not in column_map and ("market" in text or "market rent" in text):
            column_map["market_rent"] = c
        elif "charge_code" not in column_map and ("charge" in text or "charge code" in text or text == "code"):
            column_map["charge_code"] = c
        elif "amount" not in column_map and (text == "amount" or "charge amount" in text):
            column_map["amount"] = c
        elif "deposit" not in column_map and ("deposit" in text):
            column_map["deposit"] = c
        elif "move_in" not in column_map and ("move in" in text or "movein" in text):
            column_map["move_in"] = c
        elif "lease_expiration" not in column_map and ("lease expiration" in text or "expiration" in text or "lease expire" in text):
            column_map["lease_expiration"] = c
        elif "move_out" not in column_map and ("move out" in text or "moveout" in text):
            column_map["move_out"] = c
        elif "balance" not in column_map and ("balance" in text):
            column_map["balance"] = c
        elif "resident_name" not in column_map and ("resident" in text or "resident name" in text or "name" in text or "tenant" in text):
            column_map["resident_name"] = c

    # Special check: If Row 5 had "Resident" and "Name" in separate adjacent columns
    if "resident_name" not in column_map:
        for c in range(max_len - 1):
            if _key(row5[c]) == "resident" and _key(row5[c + 1]) == "name":
                column_map["resident_name"] = c
                break

    # Required logical columns
    required = ("unit", "unit_type", "sq_ft", "resident_name", "market_rent", "charge_code", "amount")
    if not all(col in column_map for col in required):
        raise _Failure("UNSUPPORTED_XLSX_LAYOUT")

    return column_map


def _parse_workbook(raw: bytes) -> dict[str, Any]:
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        entries = archive.infolist()
        if len(entries) > 1024 or sum(i.file_size for i in entries) > _MAX_BYTES:
            raise _Failure("INPUT_LIMIT_EXCEEDED")
        if any(i.flag_bits & 1 or "vbaproject" in i.filename.lower()
               or "externallinks/" in i.filename.lower()
               or "activex/" in i.filename.lower()
               or "ctrlprops/" in i.filename.lower()
               or "embeddings/" in i.filename.lower() for i in entries):
            raise _Failure("MALFORMED_INPUT")

    try:
        from openpyxl import load_workbook
    except ImportError:
        raise _Failure("XLSX_DEPENDENCY_MISSING") from None

    book = load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
    try:
        target_sheet = None
        target_sheet_idx = 1
        for idx, ws in enumerate(book.worksheets, start=1):
            # Inspect first 4 rows for matching title
            found_title = False
            for row_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
                if row_idx > 4:
                    break
                for val in row:
                    if "rent roll with lease charges" in _key(val):
                        found_title = True
                        break
                if found_title:
                    break
            if found_title:
                target_sheet = ws
                target_sheet_idx = idx
                break

        if target_sheet is None:
            raise _Failure("UNSUPPORTED_XLSX_LAYOUT")

        # Bounds verification
        total_rows = 0
        total_cells = 0
        rows_data: list[list[str]] = []
        for row in target_sheet.iter_rows(values_only=True):
            total_rows += 1
            if total_rows > _MAX_ROWS:
                raise _Failure("INPUT_LIMIT_EXCEEDED")
            if len(row) > _MAX_COLUMNS:
                raise _Failure("INPUT_LIMIT_EXCEEDED")
            total_cells += len(row)
            if total_cells > _MAX_CELLS:
                raise _Failure("INPUT_LIMIT_EXCEEDED")
            rows_data.append([_cell_text(v) for v in row])

        if len(rows_data) < 7:
            raise _Failure("UNSUPPORTED_XLSX_LAYOUT")

        # Check split headers at rows 5 and 6 (0-indexed 4 and 5)
        row5 = rows_data[4]
        row6 = rows_data[5]
        col_map = _fold_headers(row5, row6)

        # Check section header at row 7 (0-indexed 6)
        row7 = rows_data[6]
        row7_text = " ".join(filter(None, (_key(v) for v in row7)))
        if "current" not in row7_text or "resident" not in row7_text and "unit" not in row7_text:
            raise _Failure("UNSUPPORTED_XLSX_LAYOUT")

        digest = hashlib.sha256(raw).hexdigest()
        result: dict[str, Any] = {
            "version": "1.0",
            "pms_type": "realpage",
            "source_sha256": digest,
            "residential_units": [],
            "commercial_units": [],
            "counts": {s: dict.fromkeys(_COUNTS) for s in _SCOPES},
            "summary": {
                s: {"status": "absent", "reported_counts": None, "citations": []}
                for s in _SCOPES
            },
            "issues": [],
            "status": "normalized_unvalidated",
        }

        def cite(r: int, c: int) -> dict[str, Any]:
            return {
                "source_sha256": digest,
                "sheet": target_sheet_idx,
                "row": r,
                "row_end": r,
                "column": c + 1,
            }

        def issue(
            code: str,
            r: int | None = None,
            c: int = 0,
            *,
            warning: bool = False,
            observation: dict[str, Any] | None = None,
        ) -> None:
            entry: dict[str, Any] = {
                "code": code,
                "severity": "warning" if warning else "blocker",
                "citation": cite(r, c) if r is not None else None,
            }
            if observation is not None:
                entry["observation"] = observation
            result["issues"].append(entry)
            if not warning:
                result["status"] = "blocked"

        if len(book.worksheets) > 1:
            issue("NON_INVENTORY_SHEET_IGNORED", warning=True)

        units: dict[str, dict[str, Any]] = {}
        conflicts: set[str] = set()
        current_unit: dict[str, Any] | None = None
        current_unit_id: str | None = None
        inventory_valid = True

        col_unit = col_map["unit"]
        col_type = col_map["unit_type"]
        col_sqft = col_map["sq_ft"]
        col_name = col_map["resident_name"]
        col_charge = col_map["charge_code"]
        col_amount = col_map["amount"]

        for r_idx in range(7, len(rows_data)):
            row = rows_data[r_idx]
            sheet_row = r_idx + 1  # 1-indexed

            if not any(row):
                continue

            first_text = _key(row[0]) if row else ""
            row_keys = [_key(c) for c in row if c]

            # Section termination check
            if any(k in _TOTALS or "totals" in k for k in row_keys[:2]):
                break

            if any(k in _EXCLUDED_SECTIONS or "future" in k or "applicant" in k for k in row_keys[:2]):
                issue("NON_CURRENT_SECTION_EXCLUDED", sheet_row, 0, warning=True)
                break

            unit_val = row[col_unit] if col_unit < len(row) else ""
            type_val = row[col_type] if col_type < len(row) else ""
            sqft_val = row[col_sqft] if col_sqft < len(row) else ""
            name_val = row[col_name] if col_name < len(row) else ""
            charge_val = row[col_charge] if col_charge < len(row) else ""
            amount_val = row[col_amount] if col_amount < len(row) else ""

            is_continuation = False
            if not unit_val and (charge_val or amount_val):
                is_continuation = True
            elif (
                unit_val
                and current_unit_id is not None
                and unit_val == current_unit_id
                and not type_val
                and not sqft_val
                and (charge_val or amount_val)
            ):
                is_continuation = True

            if is_continuation:
                if current_unit is None:
                    inventory_valid = False
                    issue("ORPHAN_CONTINUATION", sheet_row, col_charge)
                else:
                    evidence_item: dict[str, Any] = {}
                    if charge_val:
                        evidence_item["charge"] = cite(sheet_row, col_charge)
                    current_unit["evidence"].append(evidence_item)
                    issue("CONTINUATION_EVIDENCE", sheet_row, col_charge, warning=True)
                continue

            if unit_val:
                unit_id = unit_val

                # Determine status
                name_key = _key(name_val)
                if not name_val or name_key in {"vacant", "vacant rented", "vacant unrented"}:
                    unit_status = "vacant"
                    status_cite_col = col_name
                elif name_key in {"down", "offline", "out of service"}:
                    unit_status = "down"
                    status_cite_col = col_name
                elif name_key in {"model", "admin"}:
                    unit_status = None
                    status_cite_col = col_name
                    inventory_valid = False
                    issue("UNSUPPORTED_STATUS", sheet_row, col_name)
                elif "notice" in name_key:
                    unit_status = "occupied"
                    status_cite_col = col_name
                else:
                    unit_status = "occupied"
                    status_cite_col = col_name

                # Check if use is explicitly resolved
                type_key = _key(type_val)
                resolved_use = _TYPES.get(type_key)

                evidence_primary: dict[str, Any] = {
                    "unit_id": cite(sheet_row, col_unit),
                    "status": cite(sheet_row, status_cite_col),
                }
                if type_val:
                    evidence_primary["unit_type"] = cite(sheet_row, col_type)
                if name_val and unit_status == "occupied":
                    evidence_primary["tenant_name"] = cite(sheet_row, col_name)
                if charge_val:
                    evidence_primary["charge"] = cite(sheet_row, col_charge)

                if unit_id in conflicts:
                    issue("DUPLICATE_UNIT_CONFLICT", sheet_row, col_unit)
                    current_unit = None
                    current_unit_id = None
                    continue

                existing = units.get(unit_id)
                if existing:
                    if existing["status"] != unit_status:
                        inventory_valid = False
                        conflicts.add(unit_id)
                        existing["status"] = None
                        issue("DUPLICATE_UNIT_CONFLICT", sheet_row, col_unit)
                    else:
                        issue("DUPLICATE_UNIT_EVIDENCE", sheet_row, col_unit, warning=True)
                    existing["evidence"].append(evidence_primary)
                    current_unit = existing
                    current_unit_id = unit_id
                else:
                    current_unit = {
                        "unit_id": unit_id,
                        "status": unit_status,
                        "unit_type": resolved_use,
                        "tenant_name": "[REDACTED]",
                        "evidence": [evidence_primary],
                    }
                    units[unit_id] = current_unit
                    current_unit_id = unit_id

        if not units:
            inventory_valid = False
            issue("NO_CURRENT_UNITS", 7, 0)

        # Handle unresolved use units
        has_unresolved = False
        for u in units.values():
            if u["unit_type"] is None:
                has_unresolved = True
                u_row = u["evidence"][0]["unit_id"]["row"]
                issue("UNRESOLVED_UNIT_USE", u_row, col_type, observation=u)
            elif u["unit_type"] == "residential":
                result["residential_units"].append(u)
            elif u["unit_type"] == "commercial":
                result["commercial_units"].append(u)

        if has_unresolved or not inventory_valid:
            result["status"] = "blocked"
            for scope in _SCOPES:
                result["counts"][scope] = dict.fromkeys(_COUNTS)
        else:
            for scope in _SCOPES:
                scope_units = result[f"{scope}_units"]
                derived = {
                    "occupied": sum(u["status"] == "occupied" for u in scope_units),
                    "vacant": sum(u["status"] == "vacant" for u in scope_units),
                    "down": sum(u["status"] == "down" for u in scope_units),
                    "total": len(scope_units),
                }
                result["counts"][scope] = derived

        return result

    finally:
        book.close()


def normalize_lease_charge_xlsx(raw_bytes: bytes) -> dict[str, Any]:
    """Parse and normalize split-header lease-charge XLSX into v1 observation dict.

    Never executes formulas, external links, or macros.
    Tenant identities are strictly redacted.
    """
    code: str = "MALFORMED_INPUT"
    try:
        with warnings.catch_warnings():
            warnings.simplefilter("error")
            if not isinstance(raw_bytes, bytes):
                raise _Failure("MALFORMED_INPUT")
            if len(raw_bytes) > _MAX_BYTES:
                raise _Failure("INPUT_LIMIT_EXCEEDED")
            if not raw_bytes:
                raise _Failure("EMPTY_INPUT")
            if not raw_bytes.startswith(b"PK"):
                raise _Failure("UNSUPPORTED_XLSX_FORMAT")
            return _parse_workbook(raw_bytes)
    except _Failure as exc:
        code = exc.args[0]
    except Exception:
        code = "MALFORMED_INPUT"

    raise RentRollNormalizationError(code) from None
