"""Isolated bounded UW/recon worker, launched only by slice_b's host API.

No substitute arithmetic. The existing engine owns validation, precision and
financial definitions. Existing recon owns all comparison calculations.
"""
from __future__ import annotations

import contextlib
import json
import math
import os
from decimal import Decimal
from pathlib import Path
import sys


def write(path: Path, data: dict) -> None:
    with path.open("x", encoding="utf-8") as handle:
        json.dump(data, handle, indent=2, sort_keys=True, allow_nan=False,
                  default=lambda value: str(value) if isinstance(value, Decimal) else value)


def main() -> int:
    os.umask(0o077)
    engine_root, folder = (Path(v) for v in sys.argv[1:])
    sys.path.insert(0, str(engine_root))
    stage = "imports"
    try:
        from engine.engine import run_underwriting
        import runs.build_underwriting_reconciliation as recon
        from openpyxl import load_workbook

        # The recon annualizer itself accepts partial month coverage. Do not
        # silently turn a partial source into a T12; this is an input gate only.
        stage = "evidence_validation"
        workbook = load_workbook(folder / "t12.xlsx", read_only=True, data_only=True)
        try:
            sheet = workbook[workbook.sheetnames[0]]
            months = recon._detect_month_indices(sheet)
            if len(months) != 12:
                raise ValueError("T12 requires exactly twelve recognized source month columns")
            numeric_rows = 0
            for row in sheet.iter_rows(values_only=True):
                cells = [row[i] if i < len(row) else None for i in months]
                numeric = [isinstance(v, (int, float)) and not isinstance(v, bool) for v in cells]
                if any(numeric):
                    if not all(numeric) or any(not math.isfinite(v) for v in cells):
                        raise ValueError("A T12 numeric row has an unknown/nonfinite source month; not zero")
                    numeric_rows += 1
            if not numeric_rows:
                raise ValueError("Empty T12 evidence")
        finally:
            workbook.close()
        stage = "engine"
        inputs = json.loads((folder / "inputs.json").read_bytes(), parse_float=Decimal)
        output = run_underwriting(inputs, federation_mode=True)
        write(folder / "engine_output.json", output)
        write(folder / "validation.json", output["_validator_report"])
        stage = "reconciliation"
        sys.argv = [str(engine_root / "runs/build_underwriting_reconciliation.py"),
                    "--t12", str(folder / "t12.xlsx"),
                    "--broker-json", str(folder / "broker.json"),
                    "--underwriting-json", str(folder / "engine_output.json"),
                    "--output-dir", str(folder)]
        with (folder / "recon.stdout.log").open("x") as stdout, (folder / "recon.stderr.log").open("x") as stderr:
            with contextlib.redirect_stdout(stdout), contextlib.redirect_stderr(stderr):
                recon.main()
        bridge = json.loads((folder / "underwriting_reconciliation.json").read_bytes())
        missing = [{"bucket": row["bucket"], "missing": [k for k in ("actual", "broker", "house") if row.get(k) is None]}
                   for row in bridge.get("comparison_rows", [])
                   if any(row.get(k) is None for k in ("actual", "broker", "house"))]
        # Neither a zero CLI exit nor absence of flags proves revenue coverage.
        # Preserve every null and flag from the existing recon, never zero-fill.
        error = "RECON_EVIDENCE_INCOMPLETE" if missing or not bridge.get("comparison_rows") else "RECON_REVIEW_NOT_IMPLEMENTED"
        write(folder / "worker_result.json", {
            "status": "engine_recon_draft", "error": error, "certified": False,
            "engine_executed": True, "reconciliation_executed": True,
            "recon_missing": missing, "recon_flags": bridge.get("flags", []),
            "remaining_seam": "Existing recon emits comparisons/flags, not authoritative coverage and reviewed PASS. Certification is NOT_IMPLEMENTED.",
        })
        (folder / "draft_report.md").write_text(
            "# Synthetic underwriting draft — NOT CERTIFIED\n\n"
            "Existing engine validation and reconciliation executed. No live assumptions adopted.\n"
            f"\nCertification blocker: `{error}`. See immutable manifest, engine_output.json, "
            "validation.json and underwriting_reconciliation.json for exact evidence.\n"
            "No publishing, price discovery or investment recommendation is authorized.\n")
        return 0
    except Exception as exc:
        write(folder / "worker_result.json", {
            "status": "blocked", "certified": False,
            "error": {"imports": "NOT_IMPLEMENTED", "evidence_validation": "EVIDENCE_INCOMPLETE",
                      "engine": "ENGINE_VALIDATION_OR_EXECUTION_FAILED", "reconciliation": "RECONCILE_FAILED"}[stage],
            "stage": stage, "exception_type": type(exc).__name__, "message": str(exc),
            "engine_executed": (folder / "engine_output.json").exists(), "reconciliation_executed": False,
        })
        import traceback
        traceback.print_exc()
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
