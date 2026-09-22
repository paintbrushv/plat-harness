"""Versioned synthetic execution/recon review bridge; no financial arithmetic.

The existing contracts registry is the only review authority. Host configuration
pins that registry; neither tool requests nor model output can provide it. An
explicit execution approval is required in addition to classifier/policy review.
Reviews are append-only sidecars, never upgrades to an immutable engine run.
"""
from __future__ import annotations

from decimal import Decimal, InvalidOperation
import io
import json
import math
import os
from pathlib import Path
import tempfile

from plat_harness import contracts
from plat_harness.adapters import slice_b as b

EXECUTION_VERSION = "slice-b-execution/1.0.0"
REVIEW_VERSION = "slice-b-recon-review/1.0.0"
IDENTITY = {"subject_id", "run_id", "input_sha256", "policy_version"}
REVIEW_ARTIFACTS = ("manifest.json", "inputs.json", "approval.json", "policy.json", "classification.json",
                    "t12.xlsx", "broker.json", "engine_output.json", "validation.json",
                    "underwriting_reconciliation.json", "worker_result.json", "child_exit.json")
# Existing recon's fixed comparison surface. Missing rows are missing coverage,
# not proof that an economic category is zero or inapplicable.
BUCKETS = frozenset({"gross_potential_rent", "vacancy_loss", "concessions", "employee_model", "bad_debt",
                    "loss_to_lease", "commercial_income", "utility_reimbursements", "other_income",
                    "insurance", "real_estate_taxes", "payroll", "utilities", "repairs_maintenance",
                    "make_ready", "contract_services", "marketing", "general_administrative",
                    "management_fees", "replacement_reserves", "noi"})
MINIMUM_SCOPE = frozenset({"gross_potential_rent", "vacancy_loss", "concessions", "loss_to_lease", "insurance", "real_estate_taxes"})


def exact(value, keys, where):
    contracts._object(value, set(keys), where)


def native(value):
    """contracts.py canonical JSON uses native JSON numbers, not Decimal strings."""
    if isinstance(value, Decimal):
        result = float(value)
        if not math.isfinite(result) or (value and result == 0):
            b.refuse("INVALID_CONTRACT", "Contract number exceeds finite JSON range.")
        return result
    if isinstance(value, dict):
        return {k: native(v) for k, v in value.items()}
    if isinstance(value, list):
        return [native(v) for v in value]
    return value


def decode(data):
    return native(b.decode(data))


def code_hashes():
    _, _, engine = b._engine()
    return {"engine_sha256": engine, "adapter_sha256": b.digest(Path(b.__file__).read_bytes()),
            "worker_sha256": b.digest(Path(b.__file__).with_name("slice_b_worker.py").read_bytes()),
            "review_bridge_sha256": b.digest(Path(__file__).read_bytes()),
            "contracts_sha256": b.digest(Path(contracts.__file__).read_bytes())}


def host_registry():
    path = os.environ.get("PLAT_HARNESS_CONTRACT_REGISTRY_PATH")
    pin = os.environ.get("PLAT_HARNESS_CONTRACT_REGISTRY_SHA256")
    if not path or not isinstance(pin, str) or not b.SHA.fullmatch(pin):
        b.refuse("APPROVAL_REQUIRED", "Pinned host contracts registry is required.")
    raw = b._read(path)
    if b.digest(raw) != pin:
        b.refuse("APPROVAL_REQUIRED", "Host registry bytes differ from the host pin.")
    registry = decode(raw)
    for approval_id, record in registry.items():
        if not b.ID.fullmatch(approval_id):
            b.refuse("INVALID_CONTRACT", "Exact host approval identifier required.")
        exact(record, {"actor_id", "payload_sha256", "approval_sha256"}, "host approval record")
        contracts._text(record["actor_id"], "actor_id")
        contracts._sha(record["payload_sha256"])
        contracts._sha(record["approval_sha256"])
    return registry


def approved(value, registry):
    # Reuse the established authority binding, including the entire provenance
    # record. No parallel approval registry or weaker 'reviewed=true' shortcut.
    contracts._approval(value["approval"], {k: v for k, v in value.items() if k != "approval"}, registry)


def evidence_exists(evidence, sources):
    """Resolve only named persisted inputs; a locator never grants filesystem IO."""
    for item in evidence:
        name = item["artifact"]
        if name not in sources:
            b.refuse("UNSAFE_PATH", "Evidence must name an exact persisted source artifact.")
        data = sources[name]
        if item["sha256"] != b.digest(data):
            b.refuse("HASH_MISMATCH", "Contract source evidence bytes changed.", artifact=name)
        locator = item["locator"]
        if set(locator) == {"json_pointer"} and name.endswith(".json"):
            value = decode(data)
            try:
                for token in locator["json_pointer"][1:].split("/"):
                    token = token.replace("~1", "/").replace("~0", "~")
                    if isinstance(value, list):
                        if not token.isdecimal() or str(int(token)) != token:
                            raise ValueError()
                        value = value[int(token)]
                    else:
                        value = value[token]
                if value is None:
                    raise ValueError()
            except (KeyError, IndexError, TypeError, ValueError):
                b.refuse("EVIDENCE_INCOMPLETE", "Evidence locator does not resolve to supplied nonnull evidence.")
        elif set(locator) == {"sheet", "row"} and name == "t12.xlsx":
            from openpyxl import load_workbook
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            try:
                if locator["sheet"] not in workbook.sheetnames:
                    b.refuse("EVIDENCE_INCOMPLETE", "Evidence sheet is absent.")
                sheet = workbook[locator["sheet"]]
                if locator["row"] > sheet.max_row or not any(v is not None for v in next(sheet.iter_rows(min_row=locator["row"], max_row=locator["row"], values_only=True))):
                    b.refuse("EVIDENCE_INCOMPLETE", "Evidence row is absent/empty.")
            finally:
                workbook.close()
        else:
            b.refuse("NOT_IMPLEMENTED", "This synthetic bridge resolves JSON pointers and T12 sheet/row evidence only.")


def number(value):
    try:
        if isinstance(value, bool) or value is None:
            raise ValueError()
        n = Decimal(str(value))
        if not n.is_finite() or (n and not -308 <= n.adjusted() <= 308) or not math.isfinite(float(n)):
            raise ValueError()
        return n
    except (ValueError, InvalidOperation, OverflowError):
        b.refuse("INVALID_CONTRACT", "Finite bounded decimal required.")


def execution_gates(identity, canonical, envelope, policy, policy_hash, evidence):
    envelope, policy = native(envelope), native(policy)
    exact(envelope, IDENTITY | {"contract_version", "synthetic", "intake_validated", "analysis_window", "tax_applicability",
                               "policy_sha256", "evidence_sha256", "code_sha256", "overrides", "approval"}, "execution approval")
    if envelope["contract_version"] != EXECUTION_VERSION:
        b.refuse("NOT_IMPLEMENTED", "Unsupported execution review contract.")
    if any(envelope[k] != v for k, v in identity.items()):
        b.refuse("APPROVAL_MISMATCH", "Execution review identity does not match the exact request.")
    if envelope["synthetic"] is not True or envelope["intake_validated"] is not True:
        b.refuse("REVIEW_REQUIRED", "Explicit synthetic intake approval is required.")
    if envelope["code_sha256"] != code_hashes():
        b.refuse("STALE", "Execution approval binds different engine/bridge/contracts code.")
    if envelope["overrides"] != []:
        b.refuse("POLICY_CONFLICT", "Execution overrides are not implemented; approve new canonical bytes instead.")
    registry = host_registry()
    approved(envelope, registry)
    classification = decode(evidence["classification"])
    contracts.require_approved_policy(classification, policy, approval_registry=registry)
    if any(classification[k] != identity[k] for k in ("subject_id", "run_id", "input_sha256")):
        b.refuse("APPROVAL_MISMATCH", "Classification is for different canonical bytes or subject/run.")
    if policy["version"] != identity["policy_version"] or policy_hash != envelope["policy_sha256"]:
        b.refuse("POLICY_CONFLICT", "Exact approved policy version and bytes required.")
    expected = {name: b.digest(raw) for name, raw in evidence.items()}
    if envelope["evidence_sha256"] != expected:
        b.refuse("HASH_MISMATCH", "Execution review evidence differs from supplied bytes.")
    inputs = decode(canonical)
    if envelope["analysis_window"] != inputs["time_grid"]:
        b.refuse("NEEDS_ANALYSIS_WINDOW", "Execution review must bind the exact analysis window.")
    sources = {"inputs.json": canonical, "broker.json": evidence["broker"], "t12.xlsx": evidence["t12"]}
    evidence_exists(classification["evidence"], sources)
    tax = inputs["metadata"]["property_summary"]["property_tax_policy"]
    supplied = envelope["tax_applicability"]
    tax_fields = {"asset_id", "tax_year", "parcel_ids", "unit", "millage_rate_mills"}
    exact(supplied, tax_fields, "tax applicability")
    applicability = inputs["metadata"]["property_summary"].get("tax_applicability")
    if (supplied != applicability or number(supplied["millage_rate_mills"]) != number(tax["millage_rate_mills"])
            or supplied["asset_id"] != identity["subject_id"]
            or type(supplied["tax_year"]) is not int
            or supplied["tax_year"] != int(inputs["time_grid"]["analysis_start_date"][:4])
            or supplied["unit"] != "mills_per_1000"
            or not isinstance(supplied["parcel_ids"], list) or not supplied["parcel_ids"]
            or any(not isinstance(p, str) or not b.ID.fullmatch(p) for p in supplied["parcel_ids"])
            or len(set(supplied["parcel_ids"])) != len(supplied["parcel_ids"])):
        b.refuse("TAX_APPLICABILITY_REQUIRED", "Exact asset, analysis-start tax year, unique parcels and mills_per_1000 required.")
    if tax.get("analyst_override") is not False:
        b.refuse("POLICY_CONFLICT", "Tax overrides need a separately implemented review contract; not assumed approved.")
    bindings = {"millage_rate_mills": tax["millage_rate_mills"], "interest_rate": inputs["debt_terms"]["rate"],
                "exit_cap": inputs["exit_assumptions"]["exit_cap_rate"], "price_basis": inputs["purchase_assumptions"]["purchase_price"]}
    names = set()
    for assumption in policy["assumptions"]:
        evidence_exists(assumption["evidence"], sources)
        if assumption["value"] is None:
            continue
        name = assumption["name"]
        names.add(name)
        if name not in bindings:
            b.refuse("POLICY_NOT_IMPLEMENTED", "Policy assumption has no verified canonical binding in this minimum bridge.", assumption=name)
        if number(assumption["value"]) != number(bindings[name]):
            b.refuse("POLICY_CONFLICT", "Supplied policy assumption differs from exact canonical value.", assumption=name)
    if "millage_rate_mills" not in names:
        b.refuse("MISSING_MILLAGE", "Explicit reviewed policy millage is required.")
    broker = decode(evidence["broker"])
    exact(broker, {"revenue_assumptions", "expense_assumptions", "noi"}, "synthetic broker evidence")
    for section in ("revenue_assumptions", "expense_assumptions"):
        if not isinstance(broker[section], dict) or not set(broker[section]).issubset(BUCKETS):
            b.refuse("INVALID_CONTRACT", "Unknown broker comparison field.")
        for value in broker[section].values():
            number(value)
    number(broker["noi"])
    # Compatibility view for the existing nonfinancial gates, only AFTER the
    # independently approved execution envelope and contracts have validated.
    # Nothing is written, approved, selected or applied by this projection.
    return ({**identity, "status": "approved", "synthetic": True, "reviewer": envelope["approval"]["actor_id"],
             "conflicts": [], "intake_validated": True, "analysis_window": envelope["analysis_window"],
             "classification": {"review_state": "approved", "deal_type": classification["deal_type"], "evidence_locators": classification["evidence"]},
             "policy_sha256": policy_hash, "evidence_sha256": expected},
            {"version": policy["version"], "status": "approved", "conflicts": [], "overrides": []})


def assess_reconciliation(recon, scope):
    """Coverage oracle only. Never computes, fills or rounds financial values."""
    exact(recon, {"house_year", "trailing_actuals", "broker_snapshot", "house_underwrite", "comparison_rows", "property_tax_calculation", "flags"}, "reconciliation")
    if recon["house_year"] != "1" or not isinstance(recon["flags"], list):
        b.refuse("INVALID_CONTRACT", "Only the actual first analysis-year reconciliation is supported.")
    if (not isinstance(scope, list) or any(not isinstance(v, str) or v not in BUCKETS for v in scope)
            or len(set(scope)) != len(scope) or not MINIMUM_SCOPE.issubset(scope)):
        b.refuse("INVALID_CONTRACT", "Unique explicit review scope including the minimum comparison surface required.")
    rows = recon["comparison_rows"]
    if not isinstance(rows, list):
        b.refuse("INVALID_CONTRACT", "Actual reconciliation rows required.")
    by_bucket = {}
    keys = {"bucket", "kind", "actual", "broker", "house", "broker_vs_actual_pct", "house_vs_actual_pct", "diagnosis", "diagnosis_note"}
    for row in rows:
        if not isinstance(row, dict) or not isinstance(row.get("bucket"), str) or row["bucket"] not in BUCKETS:
            b.refuse("INVALID_CONTRACT", "Unknown recon row.")
        bucket = row["bucket"]
        exact(row, keys | ({"broker_vs_actual_amount", "house_vs_actual_amount", "house_property_tax_provenance"} if bucket == "real_estate_taxes" else set()), "recon row")
        if bucket in by_bucket:
            b.refuse("INVALID_CONTRACT", "Duplicate recon bucket.")
        by_bucket[bucket] = row
        for side in ("actual", "broker", "house"):
            if row[side] is not None:
                number(row[side])
    missing = [{"bucket": bucket, "side": side} for bucket in scope for side in ("actual", "broker", "house")
               if bucket not in by_bucket or by_bucket[bucket][side] is None]
    uncovered = sorted(BUCKETS - set(scope))
    if missing:
        status, blocker = "REFUSED", "RECON_EVIDENCE_INCOMPLETE"
    elif uncovered:
        status, blocker = "SYNTHETIC_SCOPED_PASS", "RECON_COVERAGE_INCOMPLETE"
    else:
        status, blocker = "SYNTHETIC_REVIEWED_PASS", "LIVE_CERTIFICATION_NOT_IMPLEMENTED"
    return {"review_status": status, "blocker": blocker, "missing_cells": missing, "uncovered_buckets": uncovered,
            "synthetic_pass_eligible": not missing and not uncovered, "certification_eligible": False, "certified": False}


def review_run(saved):
    """Read trusted review via host path and persist an immutable review sidecar.

    The supplied mapping is only a lookup identity; all outputs are read back
    through load_run. No request registry/path/rank is accepted by this API.
    """
    identity = {k: saved[k] for k in IDENTITY}
    saved = b.load_run(**identity)
    folder = Path(saved["artifact_dir"])
    if decode(b._read(folder / "approval.json")).get("contract_version") != EXECUTION_VERSION:
        b.refuse("REVIEW_REQUIRED", "Versioned execution approval is required for reviewed eligibility.")
    path = os.environ.get("PLAT_HARNESS_SYNTHETIC_REVIEW_PATH")
    if not path:
        b.refuse("APPROVAL_REQUIRED", "Host-owned synthetic recon review path is required.")
    raw = b._read(path)
    review = decode(raw)
    exact(review, IDENTITY | {"contract_version", "review_id", "synthetic", "scope", "artifact_sha256", "code_sha256",
                             "decision", "notes", "flag_dispositions", "approval"}, "recon review")
    if review["contract_version"] != REVIEW_VERSION or review["synthetic"] is not True:
        b.refuse("LIVE_RUN_NOT_AUTHORIZED", "Only explicit synthetic review contracts are supported.")
    if not isinstance(review["review_id"], str) or not b.ID.fullmatch(review["review_id"]):
        b.refuse("INVALID_ID", "Exact review identifier required.")
    if any(review[k] != v for k, v in identity.items()):
        b.refuse("APPROVAL_MISMATCH", "Recon review belongs to different subject/run/input/policy.")
    if review["code_sha256"] != code_hashes():
        b.refuse("STALE", "Recon review code is stale.")
    if review["decision"] not in ("PASS", "REFUSE"):
        b.refuse("INVALID_CONTRACT", "Explicit PASS or REFUSE review decision required.")
    contracts._text(review["notes"], "review notes")
    approved(review, host_registry())
    expected = {name: b.digest(b._read(folder / name)) for name in REVIEW_ARTIFACTS}
    if review["artifact_sha256"] != expected:
        b.refuse("HASH_MISMATCH", "Recon review must bind every exact saved input, output, approval and manifest.")
    recon = decode(b._read(folder / "underwriting_reconciliation.json"))
    assessment = assess_reconciliation(recon, review["scope"])
    dispositions = review["flag_dispositions"]
    if not isinstance(dispositions, list):
        b.refuse("INVALID_CONTRACT", "Explicit recon flag dispositions required.")
    for disposition in dispositions:
        exact(disposition, {"sha256", "reason"}, "flag disposition")
        contracts._sha(disposition["sha256"])
        contracts._text(disposition["reason"], "flag reason")
    approved_flags = [d["sha256"] for d in dispositions]
    if len(set(approved_flags)) != len(approved_flags) or set(approved_flags) != {contracts.canonical_sha256(flag) for flag in recon["flags"]}:
        b.refuse("RECON_REVIEW_REQUIRED", "Every actual recon flag requires an exact reviewed disposition.")
    validation = decode(b._read(folder / "validation.json"))
    child = decode(b._read(folder / "child_exit.json"))
    if (review["decision"] != "PASS" or validation.get("status") != "PASS"
            or saved.get("engine_executed") is not True or saved.get("reconciliation_executed") is not True
            or type(child.get("exit_code")) is not int or child["exit_code"] != 0 or child.get("timed_out") is not False):
        assessment.update(review_status="REFUSED", blocker="RECON_REVIEW_REQUIRED", synthetic_pass_eligible=False)
    result = {**identity, **assessment, "review_id": review["review_id"], "review_sha256": b.digest(raw),
              "artifact_sha256": expected, "code_sha256": review["code_sha256"], "scope": review["scope"],
              "decision": review["decision"], "definition_owner": "existing UW engine and reconciliation; review is coverage-only"}
    # Separate append-only namespace; never edit the existing engine directory.
    parent = b.safe_path(b._root() / identity["subject_id"] / "_reviews")
    parent.mkdir(mode=0o700, exist_ok=True)
    parent = b.safe_path(parent / identity["run_id"])
    parent.mkdir(mode=0o700, exist_ok=True)
    target = b.safe_path(parent / review["review_id"])
    encoded = b.encode(result)
    if not target.exists():
        stage = Path(tempfile.mkdtemp(prefix=".review-", dir=parent))
        b._write(stage / "review.json", raw)
        b._write(stage / "result.json", encoded)
        try:
            stage.rename(target)
        except OSError:
            b.refuse("RUN_CONFLICT", "Concurrent review exists; private attempt retained.")
    if b._read(target / "review.json") != raw or b._read(target / "result.json") != encoded:
        b.refuse("RUN_CONFLICT", "Review ID already binds changed review/evidence/result bytes.")
    return {**decode(b._read(target / "result.json")), "review_artifact_dir": str(target)}
